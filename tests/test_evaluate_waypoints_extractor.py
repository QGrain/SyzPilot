import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from analyzer import evaluate_waypoints_extractor as evaluator
import waypoints_extractor as extractor_module


def make_node(func_name: str, location: str, original_index: int):
    return evaluator.NodeSnapshot(
        func_name=func_name,
        location=location,
        is_inline=False,
        original_index=original_index,
        bb_offset=1,
        bb_count=8,
        hot_entry_score=0.875,
        value_score=1.0,
    )


def make_snapshot(key: str, nodes):
    snapshot = evaluator.StageSnapshot(key=key, nodes=list(nodes))
    snapshot.resolved_targets = [node.target for node in nodes]
    snapshot.pcs64 = [f"0xffffffff0000000{index + 1}" for index in range(len(nodes))]
    snapshot.pcs32 = [evaluator.format_pc32(pc) for pc in snapshot.pcs64]
    snapshot.resolution_errors = ["" for _ in nodes]
    return snapshot


class WaypointRegressionExporterTests(unittest.TestCase):
    def setUp(self):
        self.target = make_node("target", "kernel/demo.c:30", 2)
        self.middle = make_node("middle", "kernel/demo.c:20", 1)
        self.entry = make_node("entry", "kernel/demo.c:10", 0)

    def test_ordered_subsequence_rejects_reordering(self):
        parent = [self.entry.identity, self.middle.identity, self.target.identity]
        self.assertTrue(
            evaluator.ordered_subsequence(
                [self.entry.identity, self.target.identity], parent
            )
        )
        self.assertFalse(
            evaluator.ordered_subsequence(
                [self.target.identity, self.entry.identity], parent
            )
        )

    def test_case_row_separates_display_and_fuzzer_pc_directions(self):
        nodes = [self.target, self.middle, self.entry]
        final_snapshot = make_snapshot("outlier_removal", nodes)
        result = evaluator.CaseResult(
            metadata={"ID": "1", "Title": "BUG in target"},
            status="ok",
            snapshots={"outlier_removal": final_snapshot},
        )
        row = evaluator.case_row(result)

        self.assertEqual(
            row["Final Waypoints (target->entry)"].splitlines(),
            [node.target for node in nodes],
        )
        self.assertEqual(
            evaluator.json.loads(row["Final Listed PCs (target->entry)"]),
            final_snapshot.pcs32,
        )
        self.assertEqual(
            evaluator.json.loads(row["Fuzzer target_pcs (entry->target)"]),
            list(reversed(final_snapshot.pcs32)),
        )

    def test_canonical_payload_uses_native_arrays_and_null_pcs(self):
        snapshot = make_snapshot("outlier_removal", [self.target, self.entry])
        snapshot.pcs64[1] = evaluator.ZERO_PC64
        snapshot.pcs32[1] = evaluator.ZERO_PC32
        snapshot.resolution_errors[1] = "unresolved"
        result = evaluator.CaseResult(
            metadata={
                "ID": "1",
                "Title": "BUG in target",
                "Bug Position": "kernel/demo.c:30",
            },
            status="ok",
            snapshots={"outlier_removal": snapshot},
        )

        payload = evaluator.build_canonical_payload(
            [result], {"Run ID": "test-run"}
        )
        final = payload["cases"][0]["final"]

        self.assertEqual(payload["schema_version"], "2.0")
        self.assertEqual(
            final["waypoints_target_to_entry"],
            [self.target.target],
        )
        self.assertIsNone(
            payload["cases"][0]["stages"]["outlier_removal"]["chain"][
                "pcs64_target_to_entry"
            ][1]
        )
        self.assertEqual(
            final["fuzzer_pcs32_entry_to_target"],
            list(reversed(final["pcs32_target_to_entry"])),
        )

    def test_input_state_includes_bug_position_source_outside_trace(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            kernel_dir = Path(tmpdir) / "kernel"
            (kernel_dir / "arch/x86/boot").mkdir(parents=True)
            (kernel_dir / "kernel").mkdir()
            (kernel_dir / "fs").mkdir()
            (kernel_dir / "arch/x86/boot/bzImage").write_bytes(b"kernel")
            (kernel_dir / "vmlinux").write_bytes(b"symbols")
            (kernel_dir / "kernel/trace.c").write_text(
                "int trace;\n", encoding="utf-8"
            )
            (kernel_dir / "fs/bug.c").write_text("int bug;\n", encoding="utf-8")
            title = Path(tmpdir) / "case.title"
            report = Path(tmpdir) / "case.report"
            title.write_text("BUG in target\n", encoding="utf-8")
            report.write_text("Call Trace:\n", encoding="utf-8")
            trace_node = make_node("entry", "kernel/trace.c:1", 0)
            result = evaluator.CaseResult(
                metadata={
                    "ID": "1",
                    "Kernel Dir": str(kernel_dir),
                    "Title Path": str(title),
                    "Report Path": str(report),
                    "Bug Position": "fs/bug.c:1",
                },
                bug_position_source_target="target@fs/bug.c:1",
                snapshots={"call_trace": make_snapshot("call_trace", [trace_node])},
            )

            state = evaluator.case_input_state(result)

            self.assertEqual(
                set(state["referenced_sources"]),
                {"fs/bug.c", "kernel/trace.c"},
            )

    def test_operational_final_drops_zero_and_keeps_deepest_duplicate(self):
        snapshot = make_snapshot(
            "outlier_removal", [self.target, self.middle, self.entry]
        )
        snapshot.pcs64 = [
            "0xffffffff81000005",
            evaluator.ZERO_PC64,
            "0xffffffff81000005",
        ]
        snapshot.pcs32 = ["0x81000005", evaluator.ZERO_PC32, "0x81000005"]
        operational = evaluator.operational_final_snapshot(snapshot)

        self.assertEqual(operational.targets, [self.target.target])
        self.assertEqual(operational.pcs32, ["0x81000005"])

    def test_bug_position_relations_use_target_to_entry_indices(self):
        snapshot = make_snapshot(
            "outlier_removal", [self.target, self.middle, self.entry]
        )
        self.assertEqual(
            evaluator.classify_location_relation(snapshot, "kernel/demo.c:30"),
            ("target", 0),
        )
        self.assertEqual(
            evaluator.classify_function_relation(snapshot, "middle.constprop.2"),
            ("toward_entry", 1),
        )
        self.assertEqual(
            evaluator.function_name_from_resolved_target(
                "middle@/kernel/demo.c:20"
            ),
            "middle",
        )

    def test_empty_final_chain_is_a_validation_error(self):
        result = evaluator.CaseResult(
            metadata={"ID": "1"},
            status="ok",
            snapshots={"outlier_removal": make_snapshot("outlier_removal", [])},
        )
        evaluator.validate_case(result)

        failures = {
            validation.check
            for validation in result.validations
            if validation.severity == "error" and not validation.passed
        }
        self.assertIn("final_chain_nonempty", failures)

    def test_bb_without_preceding_instrumentation_site_is_removed(self):
        waypoints = extractor_module.Waypoints(
            "target+0x10/0x20 kernel/demo.c:10", "target"
        )
        extractor = object.__new__(extractor_module.WaypointsExtractor)
        extractor.waypoints = waypoints
        extractor.load_bb_info(
            {
                "target": {
                    "/kernel/demo.c:12": ["0xffffffff81000005"],
                }
            }
        )

        self.assertEqual(extractor.waypoints.length, 0)

    def test_kasan_free_trace_stops_before_related_work_and_metadata(self):
        extractor = object.__new__(extractor_module.WaypointsExtractor)
        extractor.bug_func = "bug_target"
        lines = [
            "Call Trace:",
            "bug_target+0x1/0x2 kernel/bug.c:20",
            "entry_SYSCALL_64_after_hwframe+0x1/0x2",
            "Freed by task 12:",
            "free_leaf+0x1/0x2 mm/free.c:30",
            "free_root+0x1/0x2 kernel/work.c:40",
            "Last potentially related work creation:",
            "unrelated+0x1/0x2 kernel/other.c:50",
            "page:ffff refcount:1 mapcount:0",
        ]

        self.assertEqual(
            extractor.extract_kasan_calltrace(lines).splitlines(),
            [
                "bug_target+0x1/0x2 kernel/bug.c:20",
                "free_leaf+0x1/0x2 mm/free.c:30",
                "free_root+0x1/0x2 kernel/work.c:40",
            ],
        )
        self.assertIsNone(
            extractor_module.match_stack_frame("page:ffff refcount:1 mapcount:0")
        )

    def test_normal_trace_stops_before_allocation_section_without_task_marker(self):
        extractor = object.__new__(extractor_module.WaypointsExtractor)
        lines = [
            "Call Trace:",
            "bug_target+0x1/0x2 kernel/bug.c:20",
            "root+0x1/0x2 kernel/root.c:30",
            "Allocated by task 12:",
            "alloc_leaf+0x1/0x2 mm/alloc.c:40",
        ]

        self.assertEqual(
            extractor.extract_normal_calltrace(lines).splitlines(),
            [
                "bug_target+0x1/0x2 kernel/bug.c:20",
                "root+0x1/0x2 kernel/root.c:30",
            ],
        )

    def test_direct_target_pc_deduplication_preserves_deepest_node(self):
        targets, pcs = extractor_module.deduplicate_target_pcs(
            ["entry", "middle", "target", "invalid"],
            ["0xffffffff81000005", "0xffffffff81000015", "0xffffffff81000005", "0x0"],
        )
        self.assertEqual(targets, ["middle", "target"])
        self.assertEqual(
            pcs, ["0xffffffff81000015", "0xffffffff81000005"]
        )

    def test_raw_node_outside_core_scope_has_unattempted_pc(self):
        snapshot = evaluator.StageSnapshot(key="call_trace", nodes=[self.entry])
        evaluator.resolve_snapshot(snapshot, Path("/unused"), None, core_scope=set())

        self.assertEqual(snapshot.pcs64, [None])
        self.assertEqual(snapshot.pcs32, [None])
        self.assertIn("not attempted", snapshot.resolution_errors[0])

    def test_targeted_baseline_comparison_does_not_report_removed_cases(self):
        current = [{"ID": "1", "Status": "ok"}]
        baseline = {"1": {"ID": "1", "Status": "ok"}, "2": {"ID": "2"}}

        targeted = evaluator.build_change_rows(
            current, baseline, include_removed=False
        )
        full = evaluator.build_change_rows(current, baseline, include_removed=True)

        self.assertFalse(any(row.get("Change") == "removed" for row in targeted))
        self.assertTrue(
            any(row.get("ID") == "2" and row.get("Change") == "removed" for row in full)
        )

    def test_oracle_schema_is_validated_before_run(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "oracle.json"
            path.write_text(
                json.dumps({"schema_version": "1.0", "cases": {"1": None}}),
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "must be an object"):
                evaluator.load_oracle(path)

    def test_run_case_calls_core_phases_in_extract_order(self):
        events = []

        class FakeWaypoints:
            length = 1

            def __iter__(self):
                return iter([])

            def recover_calltrace(self):
                return "target@kernel/demo.c:1"

            def calc_waypoints_value_scores(self, *_args):
                events.append("score")

        class FakeExtractor:
            MIN_BB_COUNT = 5
            MIN_WAYPOINTS_LENGTH = 4
            MAX_WAYPOINTS_LENGTH = 10
            HOT_ENTRY_SCORE_THRESHOLD = 0.9
            VALUE_SCORE_ALPHA = 0.2
            VALUE_SCORE_BETA = 0.8
            VALUE_SCORE_GAMMA = 1.0

            def __init__(self, *_args, **_kwargs):
                events.append("init")
                self.report_type = "Normal"
                self.bug_func = "target"
                self.waypoints = FakeWaypoints()

            def trace_sanitization(self):
                events.append("sanitize")

            def load_bb_info(self, _info):
                events.append("load_bb")

            def complexity_based_filtering(self):
                events.append("complexity")

            def non_trivial_stage_separation(self):
                events.append("separation")

            def recursively_greedy_triming(self):
                events.append("greedy")

            def remove_outliers(self):
                events.append("outliers")

        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            cases_root = root / "cases"
            configs_dir = root / "configs"
            (cases_root / "case_1").mkdir(parents=True)
            configs_dir.mkdir()
            (configs_dir / "case_1.title").write_text(
                "BUG in target", encoding="utf-8"
            )
            (configs_dir / "case_1.report").write_text(
                "Call Trace:\ntarget kernel/demo.c:1", encoding="utf-8"
            )

            def fake_fast_build(*_args, **kwargs):
                events.append("fast_build")
                self.assertTrue(kwargs["allow_partial"])
                return {}

            with mock.patch.object(evaluator, "WaypointsExtractor", FakeExtractor), mock.patch.object(
                evaluator, "action_fast_build", fake_fast_build
            ), mock.patch.object(
                evaluator,
                "resolve_function_for_location",
                return_value="target@kernel/demo.c:1",
            ):
                evaluator.run_case(
                    {"ID": "1", "Bug Position": "kernel/demo.c:1"},
                    cases_root,
                    configs_dir,
                )

        self.assertEqual(
            events,
            [
                "init",
                "sanitize",
                "fast_build",
                "load_bb",
                "complexity",
                "separation",
                "score",
                "greedy",
                "outliers",
            ],
        )

    def test_static_workbook_contains_required_pipeline_sheets(self):
        nodes = [self.target, self.middle, self.entry]
        snapshots = {
            key: make_snapshot(key, nodes)
            for key, _, _ in evaluator.STAGES
        }
        result = evaluator.CaseResult(
            metadata={"ID": "1", "Title": "BUG in target"},
            title_bug_func="target",
            status="ok",
            snapshots=snapshots,
        )
        evaluator.validate_case(result)
        workbook = evaluator.build_workbook(
            [result], {"Run ID": "unit-test"}, baseline_rows={}, oracle_entries={}
        )

        self.assertEqual(
            set(workbook.sheetnames),
            {
                "Run_Metadata",
                "Cases",
                "Waypoints_Long",
                "Validation",
                "Summary",
                "Changes",
                "Manual_Audit",
                "Legend",
            },
        )
        cases_headers = [cell.value for cell in workbook["Cases"][1]]
        self.assertIn("After Trace Sanitization Waypoints (target->entry)", cases_headers)
        self.assertIn("Fuzzer target_pcs (entry->target)", cases_headers)

        second_workbook = evaluator.build_workbook(
            [result], {"Run ID": "unit-test-2"}, baseline_rows={}, oracle_entries={}
        )
        second_headers = [cell.value for cell in second_workbook["Cases"][1]]
        self.assertEqual(cases_headers, second_headers)

        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "checkpoint.xlsx"
            evaluator.save_workbook_atomic(second_workbook, output)
            reloaded = load_workbook(output, read_only=True)
            self.assertIn("Cases", reloaded.sheetnames)


if __name__ == "__main__":
    unittest.main()
