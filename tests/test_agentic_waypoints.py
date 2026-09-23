import argparse
import asyncio
import json
import os
import shutil
import subprocess
import tempfile
import time
import unittest
from types import SimpleNamespace
from unittest import mock
from pathlib import Path

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from analyzer import agentic_waypoint_scorer as scorer
from analyzer import agentic_waypoints_extractor as extractor
from analyzer import waypoint_evaluation as evaluation
from analyzer.waypoint_schema import (
    AgenticExtractionDecision,
    AgenticExtractionRecord,
    AgenticTargetResolution,
    AgenticWaypoint,
    BlindScoreRecord,
    ScoringCaseInput,
)


class AgenticWaypointSchemaTests(unittest.TestCase):
    def test_target_requires_normalized_relative_source_location(self):
        waypoint = AgenticWaypoint(
            target="target@kernel/demo.c:12",
            causal_phase="trigger",
            report_evidence="report frame",
            proxy_reason="",
        )
        self.assertEqual(waypoint.target, "target@kernel/demo.c:12")
        for target in (
            "target@/kernel/demo.c:12",
            "target@../kernel/demo.c:12",
            "target@kernel/demo.c:line",
            "target@kernel/demo.c:0",
        ):
            with self.subTest(target=target), self.assertRaises(ValidationError):
                AgenticWaypoint(
                    target=target,
                    causal_phase="trigger",
                    report_evidence="report frame",
                    proxy_reason="",
                )

    def test_extraction_status_and_chain_are_consistent(self):
        base = {
            "report_kind": "normal",
            "concurrency_class": "single_thread",
            "confidence": "high",
            "rationale": "evidence",
            "unresolved_questions": [],
        }
        with self.assertRaises(ValidationError):
            AgenticExtractionDecision(
                status="ok", waypoints_target_to_entry=[], **base
            )
        with self.assertRaises(ValidationError):
            AgenticExtractionDecision(
                status="failed",
                waypoints_target_to_entry=[
                    AgenticWaypoint(
                        target="target@kernel/demo.c:12",
                        causal_phase="trigger",
                        report_evidence="frame",
                        proxy_reason="",
                    )
                ],
                **base,
            )

    def test_ok_extraction_requires_exactly_one_configured_target(self):
        base = {
            "status": "ok",
            "report_kind": "KASAN",
            "concurrency_class": "serial",
            "confidence": "high",
            "rationale": "test",
            "unresolved_questions": [],
        }
        for phases in (("trigger",), ("configured_target", "configured_target")):
            with self.subTest(phases=phases), self.assertRaises(ValidationError):
                AgenticExtractionDecision(
                    waypoints_target_to_entry=[
                        AgenticWaypoint(
                            target=f"node{index}@kernel/demo.c:{index + 1}",
                            causal_phase=phase,
                            report_evidence="report",
                            proxy_reason="",
                        )
                        for index, phase in enumerate(phases)
                    ],
                    **base,
                )

    def test_ok_extraction_requires_configured_target_first(self):
        with self.assertRaisesRegex(
            ValidationError, "configured target first"
        ):
            AgenticExtractionDecision(
                status="ok",
                report_kind="KASAN",
                concurrency_class="cross-task lifecycle",
                confidence="high",
                waypoints_target_to_entry=[
                    AgenticWaypoint(
                        target="allocation@kernel/demo.c:1",
                        causal_phase="object_allocation",
                        report_evidence="allocation stack",
                        proxy_reason="",
                    ),
                    AgenticWaypoint(
                        target="target@kernel/demo.c:12",
                        causal_phase="configured_target",
                        report_evidence="use stack",
                        proxy_reason="",
                    ),
                ],
                rationale="test",
                unresolved_questions=[],
            )

    def test_ok_record_requires_matching_nonzero_target_resolution(self):
        decision = AgenticExtractionDecision(
            status="ok",
            report_kind="KASAN",
            concurrency_class="serial",
            confidence="high",
            waypoints_target_to_entry=[
                AgenticWaypoint(
                    target="target@kernel/demo.c:12",
                    causal_phase="configured_target",
                    report_evidence="report frame",
                    proxy_reason="",
                )
            ],
            rationale="test",
            unresolved_questions=[],
        )
        base = {
            "case_id": "1",
            "title": "BUG in target",
            "bug_position": "kernel/demo.c:12",
            "model": "model",
            "effort": "xhigh",
            "codex_sdk_version": "test",
            "thread_id": "thread",
            "decision": decision,
        }
        with self.assertRaisesRegex(
            ValidationError, "configured target resolution"
        ):
            AgenticExtractionRecord(**base)
        with self.assertRaisesRegex(ValidationError, "first waypoint"):
            AgenticExtractionRecord(
                **base,
                configured_target_resolution=AgenticTargetResolution(
                    proposed_target="other@kernel/demo.c:10",
                    resolved_target="other@kernel/demo.c:10",
                    pc64="0xffffffff81000010",
                    pc32="0x81000010",
                ),
            )
        with self.assertRaisesRegex(ValidationError, "must be nonzero"):
            AgenticTargetResolution(
                proposed_target="target@kernel/demo.c:12",
                resolved_target="target@kernel/demo.c:12",
                pc64="0x0000000000000000",
                pc32="0x00000000",
            )

    def test_scoring_request_projects_only_blind_fields(self):
        case = ScoringCaseInput.model_validate(
            {
                "case_id": "1",
                "title": "BUG in target",
                "bug_position": "kernel/demo.c:12",
                "report_path": "/evidence/case_1.report",
                "kernel_dir": "/evidence/case_1",
                "candidates": [
                    {
                        "candidate_id": "c_ffffffffffffffff",
                        "method_key": "agentic_final",
                        "waypoints_target_to_entry": ["target@kernel/demo.c:12"],
                    },
                    {
                        "candidate_id": "c_0000000000000000",
                        "method_key": "script_final",
                        "waypoints_target_to_entry": ["entry@kernel/demo.c:1"],
                    },
                ],
            }
        )
        request = scorer.scoring_request(case)
        visible = json.loads(request.split("JSON bundle:\n", 1)[1])

        self.assertEqual(
            [item["candidate_id"] for item in visible["candidates"]],
            ["c_0000000000000000", "c_ffffffffffffffff"],
        )
        self.assertNotIn("method_key", request)
        self.assertNotIn("agentic_final", request)
        self.assertEqual(
            set(visible["candidates"][0]),
            {"candidate_id", "waypoints_target_to_entry"},
        )

    def test_scoring_input_rejects_duplicate_cases(self):
        candidate = {
            "candidate_id": "c_0000000000000000",
            "method_key": "script_final",
            "waypoints_target_to_entry": ["target@kernel/demo.c:12"],
        }
        case = {
            "case_id": "1",
            "title": "title",
            "bug_position": "kernel/demo.c:12",
            "report_path": "/report",
            "kernel_dir": "/kernel",
            "candidates": [candidate],
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "input.json"
            path.write_text(json.dumps({"cases": [case, case]}), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "case_id values must be unique"):
                scorer.load_input(path)

    def test_cached_score_requires_exact_candidate_chain_snapshot(self):
        case = ScoringCaseInput.model_validate(
            {
                "case_id": "1",
                "title": "title",
                "bug_position": "kernel/demo.c:12",
                "report_path": "/evidence/case_1.report",
                "kernel_dir": "/evidence/case_1",
                "candidates": [
                    {
                        "candidate_id": "c_0000000000000001",
                        "method_key": "script:final",
                        "waypoints_target_to_entry": ["target@kernel/demo.c:12"],
                    }
                ],
            }
        )
        cached = BlindScoreRecord.model_validate(
            {
                "case_id": "1",
                "model": "model",
                "effort": "xhigh",
                "codex_sdk_version": "test",
                "thread_id": "thread",
                "candidate_method_map": {
                    "c_0000000000000001": "script:final"
                },
                "candidate_chain_map": {
                    "c_0000000000000001": ["target@kernel/demo.c:12"]
                },
                "decision": {
                    "reviews": [
                        {
                            "candidate_id": "c_0000000000000001",
                            "target_fidelity": "exact",
                            "section_fidelity": "complete",
                            "causal_coherence": "coherent",
                            "parsimony": "concise",
                            "confidence": "high",
                            "evidence": [],
                            "rationale": "test",
                            "unresolved_questions": [],
                        }
                    ]
                },
            }
        )

        self.assertTrue(scorer.cached_score_matches(cached, case, "model", "xhigh"))
        changed = case.model_copy(deep=True)
        changed.candidates[0].waypoints_target_to_entry = [
            "other@kernel/demo.c:13"
        ]
        self.assertFalse(
            scorer.cached_score_matches(cached, changed, "model", "xhigh")
        )

    def test_evidence_roots_accepts_scoring_case_objects(self):
        case = ScoringCaseInput.model_validate(
            {
                "case_id": "1",
                "title": "title",
                "bug_position": "kernel/demo.c:12",
                "report_path": "/evidence/configs/case_1.report",
                "kernel_dir": "/evidence/cases/case_1",
                "candidates": [
                    {
                        "candidate_id": "c_0000000000000001",
                        "method_key": "script:final",
                        "waypoints_target_to_entry": ["target@kernel/demo.c:12"],
                    }
                ],
            }
        )
        self.assertEqual(
            scorer.evidence_roots([case]),
            ["/evidence/cases/case_1", "/evidence/configs"],
        )


class AgenticExtractorOfflineTests(unittest.TestCase):
    def test_cached_missing_input_is_rechecked_when_files_appear(self):
        decision = AgenticExtractionDecision(
            status="missing_input",
            report_kind="unknown",
            concurrency_class="unknown",
            confidence="low",
            waypoints_target_to_entry=[],
            rationale="missing",
            unresolved_questions=[],
        )
        cached = AgenticExtractionRecord(
            case_id="1",
            title="title",
            bug_position="kernel/demo.c:12",
            model="model",
            effort="xhigh",
            codex_sdk_version="test",
            thread_id="",
            decision=decision,
        )
        row = {
            "ID": "1",
            "Title": "title",
            "Bug Position": "kernel/demo.c:12",
        }
        self.assertTrue(
            extractor.cached_extraction_matches(
                cached, row, False, "model", "xhigh"
            )
        )
        self.assertFalse(
            extractor.cached_extraction_matches(
                cached, row, True, "model", "xhigh"
            )
        )
        changed_row = {**row, "Bug Position": "kernel/demo.c:13"}
        self.assertFalse(
            extractor.cached_extraction_matches(
                cached, changed_row, False, "model", "xhigh"
            )
        )

    def test_all_missing_inputs_do_not_require_codex_sdk(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            benchmark = root / "benchmark.csv"
            benchmark.write_text(
                "ID,Title,Bug Position\n1,BUG in target,kernel/demo.c:12\n",
                encoding="utf-8",
            )
            prompt = root / "prompt.md"
            prompt.write_text("Read only.", encoding="utf-8")
            output = root / "records.jsonl"
            args = argparse.Namespace(
                benchmark=benchmark,
                cases_root=root / "cases",
                configs_dir=root / "configs",
                output=output,
                prompt=prompt,
                codex_home=root / "codex",
                resolver_python=Path(os.sys.executable),
                resolver_timeout_seconds=1.0,
                target_resolution_retries=0,
                model="unavailable-model",
                effort="xhigh",
                max_concurrency=1,
                case_ids=None,
                resume=False,
            )

            records = asyncio.run(extractor.run_all(args))

            self.assertEqual(records["1"].decision.status, "missing_input")
            self.assertEqual(records["1"].codex_sdk_version, "not_installed")
            self.assertTrue(output.is_file())

    def test_incompatible_cached_record_is_ignored(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "records.jsonl"
            output.write_text(
                json.dumps(
                    {
                        "schema_version": "2.0",
                        "method": "agentic",
                        "case_id": "1",
                        "title": "BUG in target",
                        "bug_position": "kernel/demo.c:12",
                        "cache_key": "stale",
                        "input_sha256": "stale",
                        "prompt_sha256": "stale",
                        "model": "gpt-5.6-sol",
                        "effort": "xhigh",
                        "codex_sdk_version": "0.147.0",
                        "thread_id": "stale",
                        "decision": {
                            "status": "ok",
                            "report_kind": "KASAN",
                            "concurrency_class": "lifecycle",
                            "confidence": "high",
                            "waypoints_target_to_entry": [
                                {
                                    "target": "allocation@kernel/demo.c:1",
                                    "causal_phase": "object_allocation",
                                    "report_evidence": "allocation stack",
                                    "proxy_reason": "",
                                },
                                {
                                    "target": "target@kernel/demo.c:12",
                                    "causal_phase": "configured_target",
                                    "report_evidence": "use stack",
                                    "proxy_reason": "",
                                },
                            ],
                            "rationale": "stale ordering",
                            "unresolved_questions": [],
                        },
                    }
                )
                + "\n",
                encoding="utf-8",
            )

            self.assertEqual(extractor.load_existing(output), {})

    def test_resolver_timeout_terminates_and_reaps_process(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            fake_python = Path(tmpdir) / "fake-python"
            fake_python.write_text(
                "#!/usr/bin/env python3\nimport time\ntime.sleep(30)\n",
                encoding="utf-8",
            )
            fake_python.chmod(0o755)
            started = time.monotonic()
            resolution, error, repairable = asyncio.run(
                extractor.resolve_configured_target(
                    fake_python,
                    Path(tmpdir),
                    "target@kernel/demo.c:12",
                    timeout_seconds=0.05,
                )
            )
            self.assertIsNone(resolution)
            self.assertIn("timed out", error)
            self.assertFalse(repairable)
            self.assertLess(time.monotonic() - started, 3.0)

    @unittest.skipUnless(Path("/proc").is_dir(), "requires procfs")
    def test_resolver_timeout_kills_descendant_after_leader_exit(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            child_pid_path = Path(tmpdir) / "child.pid"
            fake_python = Path(tmpdir) / "fake-python"
            fake_python.write_text(
                "#!/usr/bin/env python3\n"
                "import subprocess\n"
                "from pathlib import Path\n"
                "child = subprocess.Popen(['sleep', '30'])\n"
                f"Path({str(child_pid_path)!r}).write_text(str(child.pid))\n",
                encoding="utf-8",
            )
            fake_python.chmod(0o755)
            resolution, error, repairable = asyncio.run(
                extractor.resolve_configured_target(
                    fake_python,
                    Path(tmpdir),
                    "target@kernel/demo.c:12",
                    timeout_seconds=0.1,
                )
            )
            self.assertIsNone(resolution)
            self.assertIn("timed out", error)
            self.assertFalse(repairable)
            child_pid = int(child_pid_path.read_text())
            deadline = time.monotonic() + 2.0
            while self._process_is_live(child_pid) and time.monotonic() < deadline:
                time.sleep(0.02)
            self.assertFalse(self._process_is_live(child_pid))

    def test_resolver_subprocess_accepts_strict_nonzero_payload(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            fake_python = Path(tmpdir) / "fake-python"
            payload = {
                "status": "ok",
                "resolution": {
                    "proposed_target": "target@kernel/demo.c:12",
                    "resolved_target": "target@kernel/demo.c:12",
                    "pc64": "0xffffffff81000012",
                    "pc32": "0x81000012",
                },
            }
            fake_python.write_text(
                "#!/usr/bin/env python3\n"
                f"print({json.dumps(json.dumps(payload))})\n",
                encoding="utf-8",
            )
            fake_python.chmod(0o755)
            resolution, error, repairable = asyncio.run(
                extractor.resolve_configured_target(
                    fake_python,
                    Path(tmpdir),
                    "target@kernel/demo.c:12",
                    timeout_seconds=1.0,
                )
            )
            self.assertEqual(error, "")
            self.assertEqual(resolution.pc32, "0x81000012")
            self.assertFalse(repairable)

    def test_resolver_subprocess_reports_error_and_invalid_json(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            for name, body, expected in (
                (
                    "reported-error",
                    'print(\'{"status":"error","error":"unresolvable"}\')\nraise SystemExit(1)',
                    "unresolvable",
                ),
                ("invalid-json", 'print("not-json")', "invalid JSON"),
                (
                    "mismatched-unresolvable",
                    'print(\'{"status":"unresolvable","error":"bad target"}\')\nraise SystemExit(1)',
                    "protocol mismatch",
                ),
            ):
                with self.subTest(name=name):
                    fake_python = Path(tmpdir) / name
                    fake_python.write_text(
                        "#!/usr/bin/env python3\n" + body + "\n",
                        encoding="utf-8",
                    )
                    fake_python.chmod(0o755)
                    resolution, error, repairable = asyncio.run(
                        extractor.resolve_configured_target(
                            fake_python,
                            Path(tmpdir),
                            "target@kernel/demo.c:12",
                            timeout_seconds=1.0,
                        )
                    )
                    self.assertIsNone(resolution)
                    self.assertIn(expected, error)
                    self.assertFalse(repairable)

    def test_resolver_subprocess_marks_exit_two_unresolvable_as_repairable(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            fake_python = Path(tmpdir) / "fake-python"
            fake_python.write_text(
                "#!/usr/bin/env python3\n"
                'print(\'{"status":"unresolvable","error":"bad target"}\')\n'
                "raise SystemExit(2)\n",
                encoding="utf-8",
            )
            fake_python.chmod(0o755)
            resolution, error, repairable = asyncio.run(
                extractor.resolve_configured_target(
                    fake_python,
                    Path(tmpdir),
                    "target@kernel/demo.c:12",
                    timeout_seconds=1.0,
                )
            )
            self.assertIsNone(resolution)
            self.assertIn("bad target", error)
            self.assertTrue(repairable)

    def test_repair_loop_accepts_resolved_followup(self):
        initial = self._decision("bad@kernel/demo.c:12")
        repaired = self._decision("proxy@kernel/demo.c:13")
        resolution = AgenticTargetResolution(
            proposed_target="proxy@kernel/demo.c:13",
            resolved_target="proxy@kernel/demo.c:13",
            pc64="0xffffffff81000013",
            pc32="0x81000013",
        )
        thread = mock.AsyncMock()
        thread.run.return_value = SimpleNamespace(
            final_response=repaired.model_dump_json()
        )
        with mock.patch.object(
            extractor,
            "resolve_configured_target",
            new=mock.AsyncMock(
                side_effect=[
                    (None, "unresolvable", True),
                    (resolution, "", False),
                ]
            ),
        ):
            decision, verified = asyncio.run(
                extractor.validate_and_repair_configured_target(
                    thread,
                    initial,
                    Path("/python"),
                    Path("/kernel"),
                    1,
                    1.0,
                    None,
                    None,
                    None,
                )
            )
        self.assertEqual(decision.waypoints_target_to_entry[0].target, resolution.proposed_target)
        self.assertEqual(verified, resolution)
        thread.run.assert_awaited_once()

    def test_repair_loop_rejects_exhaustion_and_false_missing_input(self):
        decision = self._decision("bad@kernel/demo.c:12")
        with mock.patch.object(
            extractor,
            "resolve_configured_target",
            new=mock.AsyncMock(return_value=(None, "unresolvable", True)),
        ):
            with self.assertRaisesRegex(RuntimeError, "remains unresolved"):
                asyncio.run(
                    extractor.validate_and_repair_configured_target(
                        mock.AsyncMock(),
                        decision,
                        Path("/python"),
                        Path("/kernel"),
                        0,
                        1.0,
                        None,
                        None,
                        None,
                    )
                )
        missing = AgenticExtractionDecision(
            status="missing_input",
            report_kind="unknown",
            concurrency_class="unknown",
            confidence="low",
            waypoints_target_to_entry=[],
            rationale="missing",
            unresolved_questions=[],
        )
        with self.assertRaisesRegex(RuntimeError, "files exist"):
            asyncio.run(
                extractor.validate_and_repair_configured_target(
                    mock.AsyncMock(),
                    missing,
                    Path("/python"),
                    Path("/kernel"),
                    0,
                    1.0,
                    None,
                    None,
                    None,
                )
            )

    def test_repair_loop_does_not_repair_resolver_infrastructure_failure(self):
        thread = mock.AsyncMock()
        with mock.patch.object(
            extractor,
            "resolve_configured_target",
            new=mock.AsyncMock(
                return_value=(None, "resolver timed out", False)
            ),
        ):
            with self.assertRaisesRegex(RuntimeError, "infrastructure failure"):
                asyncio.run(
                    extractor.validate_and_repair_configured_target(
                        thread,
                        self._decision("target@kernel/demo.c:12"),
                        Path("/python"),
                        Path("/kernel"),
                        2,
                        1.0,
                        None,
                        None,
                        None,
                    )
                )
        thread.run.assert_not_awaited()

    @staticmethod
    def _decision(target: str) -> AgenticExtractionDecision:
        return AgenticExtractionDecision(
            status="ok",
            report_kind="KASAN",
            concurrency_class="serial",
            confidence="high",
            waypoints_target_to_entry=[
                AgenticWaypoint(
                    target=target,
                    causal_phase="configured_target",
                    report_evidence="report frame",
                    proxy_reason="",
                )
            ],
            rationale="test",
            unresolved_questions=[],
        )

    @staticmethod
    def _process_is_live(pid: int) -> bool:
        stat_path = Path("/proc") / str(pid) / "stat"
        if not stat_path.is_file():
            return False
        fields = stat_path.read_text(encoding="utf-8").split()
        return len(fields) > 2 and fields[2] != "Z"

    @unittest.skipUnless(shutil.which("bwrap"), "bubblewrap is not installed")
    def test_blind_wrapper_exposes_only_evidence_and_clean_codex_home(self):
        with (
            tempfile.TemporaryDirectory(dir=Path.home()) as tmpdir,
            tempfile.TemporaryDirectory(dir="/tmp") as auth_tmpdir,
        ):
            root = Path(tmpdir)
            source_codex = Path(auth_tmpdir) / ".codex_source"
            source_codex.mkdir()
            (source_codex / "auth.json").write_text("{}", encoding="utf-8")
            (source_codex / "history.jsonl").write_text(
                "method identity", encoding="utf-8"
            )
            evidence = root / "evidence"
            evidence.mkdir()
            (evidence / "report").write_text("report", encoding="utf-8")
            (evidence / "vmlinux").write_text("symbols", encoding="utf-8")
            analysis_cache = evidence / "SyzPilot-analysis"
            analysis_cache.mkdir()
            (analysis_cache / "leaked_static_candidate").write_text(
                "script-derived functions", encoding="utf-8"
            )
            repository = root / "repository"
            repository.mkdir()
            (repository / "method-map").write_text("secret", encoding="utf-8")
            runtime = root / "runtime"
            runtime.mkdir()
            (runtime / "codex").symlink_to("/bin/bash")
            (runtime / "codex-code-mode-host").symlink_to("/bin/true")
            command = [
                str(Path(extractor.__file__).with_name("codex_blind_wrapper.sh")),
                "-c",
                (
                    f"test -r '{evidence}/report' && "
                    f"test -r '{evidence}/vmlinux' && "
                    f"test ! -e '{analysis_cache}/leaked_static_candidate' && "
                    f"test -z \"$(find '{analysis_cache}' -mindepth 1 -print -quit)\" && "
                    f"test ! -e '{repository}/method-map' && "
                    f"test ! -e '{source_codex}/history.jsonl' && "
                    "test -r \"${CODEX_HOME}/auth.json\""
                ),
            ]
            env = {
                **os.environ,
                "CODEX_HOME": str(source_codex),
                "SYZPILOT_CODEX_REAL_BIN": str(runtime / "codex"),
                "SYZPILOT_MASKED_REPO": str(repository),
                "SYZPILOT_ALLOWED_ROOTS": str(evidence),
            }
            result = subprocess.run(
                command, env=env, capture_output=True, text=True, check=False
            )
        self.assertEqual(result.returncode, 0, result.stderr)


class WaypointEvaluationTests(unittest.TestCase):
    def test_duplicate_pc_normalization_keeps_deepest_node(self):
        chain = evaluation.normalized_chain(
            [
                "target@kernel/demo.c:30",
                "middle@kernel/demo.c:20",
                "entry@kernel/demo.c:10",
            ],
            ["trigger", "trigger_path", "syscall_entry"],
            [
                "target@kernel/demo.c:30",
                "middle@kernel/demo.c:20",
                "entry@kernel/demo.c:10",
            ],
            [
                "0xffffffff81000005",
                "0xffffffff81000015",
                "0xffffffff81000005",
            ],
            ["", "", ""],
        )

        self.assertEqual(
            chain["waypoints_target_to_entry"],
            ["target@kernel/demo.c:30", "middle@kernel/demo.c:20"],
        )
        self.assertIn("duplicate_pc32", chain["dropped_nodes"][0]["reason"])
        self.assertEqual(
            chain["fuzzer_pcs32_entry_to_target"],
            ["0x81000015", "0x81000005"],
        )
        self.assertTrue(chain["configured_target_retained"])

    def test_normalization_records_lost_configured_target(self):
        chain = evaluation.normalized_chain(
            ["target@kernel/demo.c:30", "entry@kernel/demo.c:10"],
            ["trigger", "syscall_entry"],
            ["target@kernel/demo.c:30", "entry@kernel/demo.c:10"],
            [None, "0xffffffff81000005"],
            ["unresolved", ""],
        )
        self.assertFalse(chain["configured_target_retained"])
        self.assertIsNone(chain["configured_target_pc64"])

    def test_normalization_prioritizes_nonzero_configured_target_index(self):
        chain = evaluation.normalized_chain(
            [
                "sink@kernel/demo.c:40",
                "target@kernel/demo.c:30",
                "entry@kernel/demo.c:10",
            ],
            ["", "", ""],
            [
                "sink@kernel/demo.c:40",
                "target@kernel/demo.c:30",
                "entry@kernel/demo.c:10",
            ],
            [
                "0xffffffff81000005",
                "0xffffffff81000005",
                "0xffffffff81000015",
            ],
            ["", "", ""],
            configured_target_proposed_index=1,
        )
        self.assertEqual(
            chain["waypoints_target_to_entry"],
            ["target@kernel/demo.c:30", "entry@kernel/demo.c:10"],
        )
        self.assertEqual(chain["configured_target_index_target_to_entry"], 0)

    def test_normalization_keeps_chain_when_configured_target_is_absent(self):
        chain = evaluation.normalized_chain(
            ["sink@kernel/demo.c:30", "entry@kernel/demo.c:10"],
            ["trigger", "syscall_entry"],
            ["sink@kernel/demo.c:30", "entry@kernel/demo.c:10"],
            ["0xffffffff81000005", "0xffffffff81000015"],
            ["", ""],
            configured_target_proposed_index=None,
        )

        self.assertEqual(
            chain["waypoints_target_to_entry"],
            ["sink@kernel/demo.c:30", "entry@kernel/demo.c:10"],
        )
        self.assertFalse(chain["configured_target_retained"])
        self.assertIsNone(chain["configured_target_proposed_index"])
        self.assertIsNone(chain["configured_target_waypoint"])
        self.assertIsNone(chain["configured_target_pc64"])

    def test_script_target_lookup_returns_none_when_stage_drops_bug_position(self):
        case = {
            "bug_position_resolution": {
                "function": "target",
                "pc64": "0xffffffff8100002a",
            }
        }
        chain = {
            "waypoints_target_to_entry": [
                "sink@kernel/demo.c:30",
                "entry@kernel/demo.c:10",
            ],
            "pcs64_target_to_entry": [
                "0xffffffff8100000a",
                "0xffffffff8100001a",
            ],
        }

        self.assertIsNone(
            evaluation.script_configured_target_index(case, chain)
        )

    def test_agentic_target_lookup_requires_unique_phase_at_index_zero(self):
        self.assertEqual(
            evaluation.agentic_configured_target_index(
                ["configured_target", "trigger_path", "syscall_entry"]
            ),
            0,
        )
        for phases in (
            ["trigger", "configured_target", "syscall_entry"],
            ["configured_target", "configured_target"],
        ):
            with self.subTest(phases=phases), self.assertRaises(ValueError):
                evaluation.agentic_configured_target_index(phases)

    def test_agentic_dynamic_target_bonus_requires_semantic_fidelity(self):
        chain = {"configured_target_index_target_to_entry": 0}
        exact = type("Review", (), {"target_fidelity": "exact"})()
        proxy = type("Review", (), {"target_fidelity": "resolvable_proxy"})()
        unsupported = type("Review", (), {"target_fidelity": "unsupported"})()
        for review in (exact, proxy):
            self.assertEqual(
                evaluation.dynamic_configured_target_index(
                    evaluation.AGENTIC_METHOD_KEY, chain, review
                ),
                0,
            )
        for review in (unsupported, None):
            self.assertIsNone(
                evaluation.dynamic_configured_target_index(
                    evaluation.AGENTIC_METHOD_KEY, chain, review
                )
            )
        self.assertEqual(
            evaluation.dynamic_configured_target_index(
                "script:call_trace", chain, None
            ),
            0,
        )

    def test_quality_components_are_simple_and_length_weighted(self):
        def coverage(hits, total, target_hit=False):
            return {
                "Waypoint Hit Count": hits,
                "Waypoint Total": total,
                "Target Hit": target_hit,
            }

        self.assertEqual(evaluation.hit_quality_score(coverage(0, 4)), 50.0)
        self.assertEqual(evaluation.hit_quality_score(coverage(1, 4)), 60.0)
        self.assertEqual(
            evaluation.hit_quality_score(coverage(3, 3, True)), 97.5
        )
        self.assertEqual(
            evaluation.hit_quality_score(coverage(10, 10, True)), 100.0
        )
        self.assertEqual(evaluation.effective_length_score(10), 40.0)
        self.assertEqual(evaluation.effective_length_score(7), 57.1)
        self.assertEqual(evaluation.effective_length_score(4), 100.0)
        self.assertEqual(evaluation.effective_length_score(2), 100.0)
        self.assertEqual(evaluation.effective_length_score(12), 33.3)
        self.assertEqual(
            evaluation.combined_quality_score(80.0, 70.0, 90.0),
            83.0,
        )
        custom = evaluation.QualityWeights(0.2, 0.3, 0.5)
        self.assertEqual(
            evaluation.combined_quality_score(80.0, 70.0, 90.0, custom),
            82.0,
        )
        self.assertEqual(
            custom.formula(), "0.2*Semantic + 0.3*Hit + 0.5*Length"
        )
        for values in ((-0.1, 0.6, 0.5), (0.2, 0.3, 0.4)):
            with self.subTest(values=values), self.assertRaises(ValueError):
                evaluation.QualityWeights(*values)

    def test_static_and_blind_input_schema_contracts_are_independent(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            static_path = Path(tmpdir) / "static.json"
            static = {
                "schema_version": "2.0",
                "artifact_type": "script_waypoint_extraction",
                "run_metadata": {"Run ID": "run"},
                "cases": [],
            }
            static_path.write_text(json.dumps(static), encoding="utf-8")
            self.assertEqual(evaluation.load_static(static_path), static)

        scoring_input = {
            "schema_version": "2.0",
            "artifact_type": "blind_waypoint_scoring_input",
            "static_run_id": "run",
            "cases": [],
        }
        evaluation.validate_scoring_bundle(static, scoring_input, {}, {})
        scoring_input["schema_version"] = evaluation.SCHEMA_VERSION
        with self.assertRaisesRegex(ValueError, "blind scoring input schema"):
            evaluation.validate_scoring_bundle(static, scoring_input, {}, {})

    def test_rule_summary_reports_compact_quality_statistics(self):
        def item(
            case_id,
            method_key,
            stage,
            proposed_length,
            resolved_length,
            score,
            dynamic,
            hits,
        ):
            return {
                "case_id": case_id,
                "method_key": method_key,
                "extraction_method": "script",
                "stage": stage,
                "proposed_waypoints_target_to_entry": [
                    f"node_{index}@kernel/demo.c:{index + 1}"
                    for index in range(proposed_length)
                ],
                "chain": {
                    "proposed_length": proposed_length,
                    "resolved_unique_length": resolved_length,
                    "configured_target_retained": True,
                    "waypoints_target_to_entry": [
                        f"node_{index}@kernel/demo.c:{index + 1}"
                        for index in range(resolved_length)
                    ],
                    "pcs64_target_to_entry": [
                        f"0xffffffff8100{index:04x}"
                        for index in range(resolved_length)
                    ],
                },
                "agentic_quality_score": score,
                "weighted_quality_score": 0.4 * score + 0.6 * dynamic,
                "coverage": {
                    "Dynamic Coverage Score": dynamic,
                    "Waypoint Hit Ratio": hits / resolved_length,
                    "Waypoint Hit Count": hits,
                    "Waypoint Total": resolved_length,
                    "Target Hit": hits > 0,
                },
            }

        payload = {
            "schema_version": "2.0",
            "content_id": "legacy-content-digest",
            "score_semantics": {},
            "evaluations": [
                item(
                    "1", "script:call_trace", "call_trace", 20, 10, 60.0, 40.0, 2
                ),
                item(
                    "2", "script:call_trace", "call_trace", 20, 10, 70.0, 50.0, 0
                ),
                item(
                    "1",
                    "script:outlier_removal",
                    "outlier_removal",
                    11,
                    10,
                    65.0,
                    50.0,
                    4,
                ),
                item(
                    "2",
                    "script:outlier_removal",
                    "outlier_removal",
                    11,
                    10,
                    65.0,
                    40.0,
                    0,
                ),
            ],
        }
        evaluation.enrich_reporting_metrics(payload)
        self.assertEqual(payload["schema_version"], evaluation.SCHEMA_VERSION)
        self.assertNotIn("content_id", payload)
        rows = {row["Method Key"]: row for row in evaluation.summary_rows(payload)}
        raw = rows["script:call_trace"]
        final = rows["script:outlier_removal"]
        self.assertEqual(raw["Mean Candidate Length"], 20.0)
        self.assertEqual(final["Mean Candidate Length"], 11.0)
        self.assertEqual(final["Candidate Compression vs Raw"], 0.45)
        self.assertEqual(raw["Mean Operational Label Length"], 10.0)
        self.assertEqual(final["Mean Operational Label Length"], 10.0)
        self.assertEqual(final["Operational Compression vs Raw"], 0.0)
        self.assertEqual(final["Changed vs Previous N"], 2)
        self.assertEqual(final["Coverage Case Count"], 2)
        self.assertEqual(final["Overall Waypoint Hit Rate"], 0.2)
        self.assertGreater(final["Mean Quality Score"], raw["Mean Quality Score"])
        self.assertNotIn("weighted_quality_score", payload["evaluations"][0])
        self.assertIn("semantic_quality_score", payload["evaluations"][0])

    def test_scoring_input_contains_all_script_stages_and_agentic_final(self):
        stages = {
            key: {
                "chain": {
                    "waypoints_target_to_entry": [
                        f"{key}@kernel/demo.c:12"
                    ]
                }
            }
            for key, _, _ in evaluation.STAGES
        }
        static = {
            "run_metadata": {"Run ID": "run"},
            "cases": [
                {
                    "case_id": "1",
                    "status": "ok",
                    "title": "BUG in target",
                    "bug_position": "kernel/demo.c:12",
                    "paths": {
                        "report": "/evidence/case_1.report",
                        "kernel_dir": "/evidence/case_1",
                    },
                    "stages": stages,
                }
            ],
        }
        decision = AgenticExtractionDecision(
            status="ok",
            report_kind="normal",
            concurrency_class="single_thread",
            confidence="high",
            waypoints_target_to_entry=[
                AgenticWaypoint(
                    target="agent@kernel/demo.c:13",
                    causal_phase="configured_target",
                    report_evidence="frame",
                    proxy_reason="",
                )
            ],
            rationale="evidence",
            unresolved_questions=[],
        )
        record = AgenticExtractionRecord(
            case_id="1",
            title="BUG in target",
            bug_position="kernel/demo.c:12",
            model="model",
            effort="xhigh",
            codex_sdk_version="test",
            thread_id="thread",
            decision=decision,
            configured_target_resolution=AgenticTargetResolution(
                proposed_target="agent@kernel/demo.c:13",
                resolved_target="agent@kernel/demo.c:13",
                pc64="0xffffffff81000013",
                pc32="0x81000013",
            ),
        )

        payload = evaluation.build_scoring_input(static, {"1": record}, "seed")
        methods = {
            candidate["method_key"] for candidate in payload["cases"][0]["candidates"]
        }

        self.assertEqual(len(methods), len(evaluation.STAGES) + 1)
        self.assertIn(evaluation.AGENTIC_METHOD_KEY, methods)
        self.assertEqual(payload["artifact_type"], "blind_waypoint_scoring_input")

    def test_compact_workbook_projects_schema_and_derives_method_key(self):
        workbook = Workbook()
        metadata = workbook.active
        metadata.title = "Run_Metadata"
        metadata.append(["Key", "Value"])
        metadata.append(["Run ID", "run"])
        waypoints = workbook.create_sheet("Evaluation_Waypoints")
        waypoints.append(
            ["Case ID", "Extraction Method", "Stage", "Waypoint", "PC64"]
        )
        waypoints.append(["1", "script", "call_trace", "a@b.c:1", "0x1"])
        coverage = workbook.create_sheet("Coverage_Runs")
        coverage.append(["errors", "case_id", "run_idx"])
        coverage.append(["", 1, 2])
        workbook.create_sheet("Manual_Audit")

        evaluation.compact_workbook(workbook)

        self.assertNotIn("Manual_Audit", workbook.sheetnames)
        self.assertEqual(workbook["Run_Metadata"].max_row, 3)
        for sheet_name in ("Run_Metadata", "Evaluation_Waypoints", "Coverage_Runs"):
            headers = tuple(
                cell.value for cell in workbook[sheet_name][1]
            )
            self.assertEqual(headers, evaluation.PUBLISHED_WORKBOOK_COLUMNS[sheet_name])
        self.assertEqual(
            workbook["Evaluation_Waypoints"].cell(2, 2).value,
            "script:call_trace",
        )
        self.assertEqual(workbook["Coverage_Runs"].cell(2, 1).value, 1)
        self.assertEqual(workbook["Coverage_Runs"].cell(2, 2).value, 2)
        self.assertEqual(evaluation.workbook_agentic_schema(workbook), "legacy")
        self.assertIn(
                "## Agentic Legacy Caveat",
            evaluation.workbook_readme_text(workbook, "legacy.xlsx"),
        )

    def test_compact_copy_does_not_modify_canonical_manifest(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source = root / "old.xlsx"
            destination = root / "compact-copy.xlsx"
            workbook = Workbook()
            workbook.active.title = "Run_Metadata"
            workbook.active.append(["Key", "Value"])
            workbook.active.append(["Run ID", "legacy"])
            workbook.save(source)
            manifest_path = root / "manifest.json"
            manifest = {
                "artifacts": {
                    "evaluation_workbook": {
                        "path": str((root / "waypoints_evaluation.xlsx").resolve()),
                        "size": 123,
                        "sha256": "legacy",
                    }
                }
            }
            manifest_path.write_text(
                json.dumps(manifest, indent=2) + "\n", encoding="utf-8"
            )
            original = manifest_path.read_bytes()

            evaluation.compact_command(
                argparse.Namespace(input=source, output=destination)
            )

            self.assertTrue(destination.is_file())
            self.assertEqual(manifest_path.read_bytes(), original)

    def test_reporting_refresh_validates_identity_and_preserves_other_sheets(self):
        workbook = Workbook()
        metadata = workbook.active
        metadata.title = "Run_Metadata"
        metadata.append(["Key", "Value"])
        metadata.append(["Run ID", "run"])
        cases = workbook.create_sheet("Cases")
        cases.append(["ID", "Custom Case Column"])
        cases.append(["1", "keep-me"])
        current = workbook.create_sheet("Evaluation")
        current.append(["ID", "Method Key", "Custom Evaluation Column"])
        current.append(["1", "script:call_trace", "replace-me"])
        rule = workbook.create_sheet("Rule_Level_Summary")
        rule.append(["Method Key"])
        sentinel = workbook.create_sheet("Local_Audit")
        sentinel.append(["sentinel"])
        payload = {
            "provenance": {"static_run_id": "run"},
            "agentic_extractions": [{"case_id": "1"}],
            "evaluations": [
                {"case_id": "1", "method_key": "script:call_trace"}
            ],
        }

        evaluation.validate_reporting_refresh_inputs(workbook, payload)
        mismatched = {**payload, "provenance": {"static_run_id": "other"}}
        with self.assertRaisesRegex(ValueError, "different Run IDs"):
            evaluation.validate_reporting_refresh_inputs(workbook, mismatched)

        evaluation_index = workbook.sheetnames.index("Evaluation")
        workbook.remove(workbook["Evaluation"])
        evaluation.replace_reporting_sheet(
            workbook,
            "Evaluation",
            [{"ID": "1", "Method Key": "script:call_trace"}],
            evaluation_index,
        )
        self.assertEqual(workbook["Cases"]["B2"].value, "keep-me")
        self.assertEqual(workbook["Local_Audit"]["A1"].value, "sentinel")
        self.assertEqual(
            tuple(cell.value for cell in workbook["Evaluation"][1]),
            evaluation.PUBLISHED_WORKBOOK_COLUMNS["Evaluation"],
        )

    def test_manifest_artifact_path_must_match_before_refresh_update(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            expected = root / "evaluation.json"
            self.assertTrue(
                evaluation.manifest_artifact_matches_path(
                    {"path": str(expected)}, expected
                )
            )
            self.assertFalse(
                evaluation.manifest_artifact_matches_path(
                    {"path": str(root / "other.json")}, expected
                )
            )

    def test_refresh_reporting_command_updates_only_reporting_artifacts(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            json_path = root / "evaluation.json"
            workbook_path = root / "evaluation.xlsx"
            payload = {
                "schema_version": "2.0",
                "artifact_type": "waypoint_quality_evaluation",
                "content_id": "legacy-digest",
                "score_semantics": {},
                "provenance": {"static_run_id": "run"},
                "agentic_extractions": [{"case_id": "1"}],
                "coverage_runs": [],
                "evaluations": [self._reporting_item()],
            }
            json_path.write_text(json.dumps(payload), encoding="utf-8")
            workbook = Workbook()
            metadata = workbook.active
            metadata.title = "Run_Metadata"
            metadata.append(["Key", "Value"])
            metadata.append(["Run ID", "run"])
            metadata.append(["Agentic Weight", 0.4])
            metadata.append(["Dynamic Weight", 0.6])
            metadata.append(["Weighted Score Role", "legacy"])
            metadata.append(["Positive-Evidence Sensitivity Formula", "legacy"])
            cases = workbook.create_sheet("Cases")
            cases.append(["ID", "Custom Case Column"])
            cases.append(["1", "keep-me"])
            current = workbook.create_sheet("Evaluation")
            current.append(["ID", "Method Key"])
            current.append(["1", "script:call_trace"])
            rule = workbook.create_sheet("Rule_Level_Summary")
            rule.append(["Method Key"])
            sentinel = workbook.create_sheet("Local_Audit")
            sentinel.append(["sentinel"])
            original_order = list(workbook.sheetnames)
            workbook.save(workbook_path)
            manifest_path = root / "manifest.json"
            manifest_path.write_text(
                json.dumps(
                    {
                        "artifacts": {
                            "evaluation_json": {
                                "path": str(json_path),
                                "size": 0,
                            },
                            "evaluation_workbook": {
                                "path": str(workbook_path),
                                "size": 0,
                            },
                        }
                    }
                ),
                encoding="utf-8",
            )

            original_json = json_path.read_bytes()
            original_workbook = workbook_path.read_bytes()
            source_alias = root / "source-alias"
            source_alias.symlink_to(root, target_is_directory=True)
            with self.assertRaisesRegex(
                ValueError, "must not overwrite source"
            ):
                evaluation.refresh_reporting_command(
                    argparse.Namespace(
                        json=json_path,
                        workbook=workbook_path,
                        output_dir=source_alias,
                        semantic_weight=0.2,
                        hit_weight=0.3,
                        length_weight=0.5,
                    )
                )
            self.assertEqual(json_path.read_bytes(), original_json)
            self.assertEqual(workbook_path.read_bytes(), original_workbook)

            evaluation.refresh_reporting_command(
                argparse.Namespace(
                    json=json_path,
                    workbook=workbook_path,
                    output_dir=None,
                    semantic_weight=0.2,
                    hit_weight=0.3,
                    length_weight=0.5,
                )
            )

            refreshed_json = json.loads(json_path.read_text(encoding="utf-8"))
            refreshed = load_workbook(workbook_path)
            refreshed_manifest = json.loads(
                manifest_path.read_text(encoding="utf-8")
            )
            self.assertEqual(refreshed_json["schema_version"], "2.4")
            self.assertEqual(
                refreshed_json["quality_weights"],
                {
                    "semantic": 0.2,
                    "hit_quality": 0.3,
                    "effective_length": 0.5,
                },
            )
            self.assertNotIn("content_id", refreshed_json)
            self.assertEqual(refreshed.sheetnames, original_order)
            self.assertEqual(refreshed["Cases"]["B2"].value, "keep-me")
            self.assertEqual(refreshed["Local_Audit"]["A1"].value, "sentinel")
            refreshed_metadata = dict(
                refreshed["Run_Metadata"].iter_rows(min_row=2, values_only=True)
            )
            self.assertEqual(refreshed_metadata["Semantic Weight"], 0.2)
            self.assertEqual(refreshed_metadata["Hit Quality Weight"], 0.3)
            self.assertEqual(refreshed_metadata["Effective Length Weight"], 0.5)
            self.assertEqual(
                refreshed_metadata["Quality Formula"],
                "0.2*Semantic + 0.3*Hit + 0.5*Length",
            )
            self.assertNotIn("Agentic Weight", refreshed_metadata)
            self.assertNotIn("Dynamic Weight", refreshed_metadata)
            self.assertNotIn("Weighted Score Role", refreshed_metadata)
            self.assertNotIn(
                "Positive-Evidence Sensitivity Formula", refreshed_metadata
            )
            self.assertEqual(
                tuple(cell.value for cell in refreshed["Evaluation"][1]),
                evaluation.PUBLISHED_WORKBOOK_COLUMNS["Evaluation"],
            )
            self.assertEqual(refreshed["Rule_Level_Summary"].max_row, 2)
            self.assertEqual(
                refreshed_manifest["artifacts"]["evaluation_json"]["size"],
                json_path.stat().st_size,
            )
            self.assertEqual(
                refreshed_manifest["artifacts"]["evaluation_workbook"]["size"],
                workbook_path.stat().st_size,
            )
            self.assertTrue((root / "README.md").is_file())

    @staticmethod
    def _reporting_item():
        return {
            "case_id": "1",
            "title": "title",
            "candidate_id": "candidate",
            "method_key": "script:call_trace",
            "extraction_method": "script",
            "stage": "call_trace",
            "chain": {
                "proposed_length": 1,
                "resolved_unique_length": 1,
                "configured_target_retained": True,
                "configured_target_waypoint": "target@kernel/demo.c:1",
                "configured_target_pc64": "0xffffffff81000005",
                "waypoints_target_to_entry": ["target@kernel/demo.c:1"],
                "causal_phases_target_to_entry": ["configured_target"],
                "pcs64_target_to_entry": ["0xffffffff81000005"],
                "fuzzer_pcs32_entry_to_target": ["0x81000005"],
                "dropped_nodes": [],
            },
            "semantic_judgment": {
                "target_fidelity": "exact",
                "section_fidelity": "complete",
                "causal_coherence": "coherent",
                "parsimony": "concise",
                "confidence": "high",
                "evidence": [],
                "rationale": "test",
            },
            "observability_score": 15.0,
            "agentic_quality_score": 90.0,
            "weighted_quality_score": 66.0,
            "coverage": {
                "PoC Evaluation Status": "EVALUATED",
                "Waypoint Hit Count": 1,
                "Waypoint Total": 1,
                "Waypoint Hit Ratio": 1.0,
                "Target Hit": True,
                "Dynamic Coverage Score": 50.0,
                "Dynamic Evidence Confidence": "lower-bound evidence",
                "Artifact Compatibility Errors": "",
            },
        }


if __name__ == "__main__":
    unittest.main()
