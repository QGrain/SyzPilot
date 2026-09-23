#!/usr/bin/env python3
"""Run the canonical script, agentic, blind-score, and KCOV evaluation pipeline."""

from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from analyzer.evaluate_waypoints_extractor import (  # noqa: E402
    STAGES,
    file_state,
    get_vmlinux_cache_identity,
)
from analyzer.waypoint_evaluation import (  # noqa: E402
    PUBLISHED_WORKBOOK_COLUMNS,
    add_quality_weight_arguments,
    quality_weights_from_args,
    workbook_readme_text,
)


ANALYZER_DIR = PROJECT_ROOT / "analyzer"
DEFAULT_SYZPILOT_PYTHON = Path.home() / "miniconda3/envs/syzpilot/bin/python"
DEFAULT_AGENT_PYTHON = Path.home() / "miniconda3/envs/syzpilot-agent/bin/python"
DEFAULT_COVERAGE_RESULTS = (
    Path.home()
    / "kernels/SyzPilot-experiments/repro_results/waypoint_kcov_full_20260828/results.jsonl"
)


def write_json_atomic(path: Path, payload: dict[str, Any]) -> None:
    temporary = path.with_name(f".{path.name}.writing")
    temporary.write_text(
        json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    os.replace(temporary, path)


def get_git_state() -> tuple[str, bool]:
    try:
        sha = subprocess.check_output(
            ["git", "rev-parse", "HEAD"], cwd=PROJECT_ROOT, text=True
        ).strip()
        dirty = bool(
            subprocess.check_output(
                ["git", "status", "--porcelain"], cwd=PROJECT_ROOT, text=True
            ).strip()
        )
        return sha, dirty
    except (OSError, subprocess.CalledProcessError):
        return "unknown", True


def run_logged(command: list[str], log_path: Path) -> int:
    log_path.parent.mkdir(parents=True, exist_ok=True)
    with log_path.open("a", encoding="utf-8") as log:
        log.write("\n$ " + " ".join(command) + "\n")
        log.flush()
        process = subprocess.Popen(
            command,
            cwd=PROJECT_ROOT,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            bufsize=1,
        )
        assert process.stdout is not None
        for line in process.stdout:
            sys.stdout.write(line)
            sys.stdout.flush()
            log.write(line)
        process.stdout.close()
        return process.wait()


def artifact_state(paths: dict[str, Path]) -> dict[str, Any]:
    return {
        name: {
            "path": str(path.resolve()),
            "exists": path.is_file(),
            "size": path.stat().st_size if path.is_file() else None,
        }
        for name, path in paths.items()
    }


def static_artifact_matches_inputs(static_json: Path) -> bool:
    try:
        payload = json.loads(static_json.read_text(encoding="utf-8"))
        if payload.get("artifact_type") != "script_waypoint_extraction":
            return False
        for case in payload["cases"]:
            paths = case["paths"]
            kernel_dir = Path(paths["kernel_dir"])
            recorded = case["input_state"]
            referenced_sources = recorded.get("referenced_sources", {})
            current = {
                "title": file_state(Path(paths["title"])),
                "report": file_state(Path(paths["report"])),
                "bzimage": file_state(kernel_dir / "arch/x86/boot/bzImage"),
                "vmlinux_identity": get_vmlinux_cache_identity(str(kernel_dir)),
                "referenced_sources": {
                    relative: file_state(kernel_dir / relative)
                    for relative in sorted(referenced_sources)
                },
            }
            if current != recorded:
                return False
    except (OSError, KeyError, TypeError, ValueError, json.JSONDecodeError):
        return False
    return True


def validate_static_artifact_for_pipeline(static_json: Path) -> None:
    """Reject correctness failures while retaining declared missing inputs."""
    payload = json.loads(static_json.read_text(encoding="utf-8"))
    if payload.get("artifact_type") != "script_waypoint_extraction":
        raise ValueError("input is not a script waypoint extraction artifact")
    for case in payload.get("cases", []):
        case_id = str(case.get("case_id", ""))
        status = case.get("status")
        failed_errors = [
            item.get("check", "unknown")
            for item in case.get("validations", [])
            if item.get("severity") == "error" and not item.get("passed")
        ]
        if status == "missing_input":
            continue
        if status != "ok":
            raise ValueError(f"case {case_id} static extraction status is {status}")
        if failed_errors:
            raise ValueError(
                f"case {case_id} has correctness validation errors: {failed_errors}"
            )
        missing_stages = [
            stage_key
            for stage_key, _, _ in STAGES
            if not (
                case.get("stages", {}).get(stage_key, {}).get("chain", {}).get(
                    "waypoints_target_to_entry"
                )
            )
        ]
        if missing_stages:
            raise ValueError(
                f"case {case_id} has empty static stages: {missing_stages}"
            )


def json_artifact_is_readable(path: Path) -> bool:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
        return isinstance(payload, dict) and bool(payload.get("artifact_type"))
    except (OSError, TypeError, ValueError, json.JSONDecodeError):
        return False


def evaluation_artifacts_are_readable(
    evaluation_json: Path, workbook: Path, readme: Path
) -> bool:
    if not json_artifact_is_readable(evaluation_json) or not readme.is_file():
        return False
    try:
        from openpyxl import load_workbook

        payload = json.loads(evaluation_json.read_text(encoding="utf-8"))
        loaded = load_workbook(workbook, read_only=True)
        if loaded.sheetnames != list(PUBLISHED_WORKBOOK_COLUMNS):
            return False
        for sheet_name, expected_headers in PUBLISHED_WORKBOOK_COLUMNS.items():
            actual_headers = [cell.value for cell in loaded[sheet_name][1]]
            if len(actual_headers) != len(set(actual_headers)):
                return False
            if tuple(actual_headers) != tuple(expected_headers):
                return False

        expected_rows = {
            "Cases": len(payload.get("agentic_extractions", [])) + 1,
            "Agentic_Extraction": len(payload.get("agentic_extractions", [])) + 1,
            "Evaluation": len(payload.get("evaluations", [])) + 1,
            "Evaluation_Waypoints": sum(
                len(item.get("chain", {}).get("nodes", []))
                for item in payload.get("evaluations", [])
            )
            + 1,
            "Coverage_Runs": len(payload.get("coverage_runs", [])) + 1,
        }
        for sheet_name, row_count in expected_rows.items():
            if loaded[sheet_name].max_row != row_count:
                return False
        expected_readme = workbook_readme_text(loaded, workbook.name)
        return readme.read_text(encoding="utf-8") == expected_readme
    except (BadZipFile, OSError, TypeError, ValueError, json.JSONDecodeError):
        return False


def execute_step(
    name: str,
    command: list[str],
    outputs: list[Path],
    manifest: dict[str, Any],
    manifest_path: Path,
    log_path: Path,
    input_paths: list[Path] | None = None,
    input_values: dict[str, Any] | None = None,
    cleanup_paths: list[Path] | None = None,
    reuse_validator=None,
    allow_nonzero_with_outputs: bool = False,
) -> None:
    inputs = {
        str(path.resolve()): file_state(path) if path.is_file() else None
        for path in (input_paths or [])
    }
    step_identity = {
        "command": command,
        "inputs": inputs,
        "input_values": input_values or {},
    }
    previous = manifest["steps"].get(name, {})
    output_states = {
        str(path.resolve()): file_state(path) for path in outputs if path.is_file()
    }
    reusable = (
        all(path.is_file() for path in outputs)
        and previous.get("step_identity") == step_identity
        and previous.get("output_state") == output_states
        and (reuse_validator is None or reuse_validator())
    )
    if reusable:
        manifest["steps"][name] = {
            "status": "reused",
            "returncode": 0,
            "command": command,
            "step_identity": step_identity,
            "output_state": output_states,
        }
        write_json_atomic(manifest_path, manifest)
        print(f"[{name}] reusing existing artifacts", flush=True)
        return
    for output in [*outputs, *(cleanup_paths or [])]:
        if output.exists() and output.is_file():
            output.unlink()
    returncode = run_logged(command, log_path)
    succeeded = returncode == 0 or (
        allow_nonzero_with_outputs and all(path.is_file() for path in outputs)
    )
    manifest["steps"][name] = {
        "status": "complete" if succeeded else "failed",
        "returncode": returncode,
        "command": command,
        "step_identity": step_identity,
        "output_state": {
            str(path.resolve()): file_state(path)
            for path in outputs
            if path.is_file()
        },
    }
    write_json_atomic(manifest_path, manifest)
    if not succeeded:
        raise RuntimeError(f"step {name} failed with return code {returncode}")


def execute_agent_step(
    name: str,
    command: list[str],
    output: Path,
    retries: int,
    manifest: dict[str, Any],
    manifest_path: Path,
    log_path: Path,
    input_paths: list[Path] | None = None,
    input_values: dict[str, Any] | None = None,
) -> None:
    step_identity = {
        "command": command,
        "inputs": {
            str(path.resolve()): file_state(path) if path.is_file() else None
            for path in (input_paths or [])
        },
        "input_values": input_values or {},
    }
    previous = manifest["steps"].get(name, {})
    previous_status = previous.get("status")
    current_output_state = file_state(output)
    reusable = (
        output.is_file()
        and previous_status in {"complete", "reused"}
        and previous.get("step_identity") == step_identity
        and previous.get("output_state") == current_output_state
    )
    if reusable:
        manifest["steps"][name] = {
            "status": "reused",
            "attempts": previous.get("attempts", []),
            "command": command,
            "step_identity": step_identity,
            "output_state": current_output_state,
        }
        write_json_atomic(manifest_path, manifest)
        print(f"[{name}] reusing existing agent records", flush=True)
        return
    resume_interrupted = (
        previous_status in {"running", "retrying"}
        and previous.get("step_identity") == step_identity
    )
    if output.is_file() and not resume_interrupted:
        output.unlink()
        current_output_state = None
    attempts = (
        list(previous.get("attempts", []))
        if resume_interrupted
        else []
    )
    for attempt in attempts:
        if attempt.get("status") == "running":
            attempt["status"] = "interrupted"
    failed_attempts = sum(
        attempt.get("returncode") not in (None, 0) for attempt in attempts
    )
    remaining_attempts = retries + 1 - failed_attempts
    next_attempt = max(
        (int(attempt.get("attempt", 0)) for attempt in attempts),
        default=0,
    ) + 1
    for _ in range(max(remaining_attempts, 0)):
        current = list(command)
        if output.is_file() and (resume_interrupted or attempts) and "--resume" not in current:
            current.append("--resume")
        attempts.append({"attempt": next_attempt, "status": "running"})
        next_attempt += 1
        manifest["steps"][name] = {
            "status": "running",
            "attempts": attempts,
            "command": current,
            "step_identity": step_identity,
        }
        write_json_atomic(manifest_path, manifest)
        returncode = run_logged(current, log_path)
        attempts[-1] = {
            "attempt": attempts[-1]["attempt"],
            "returncode": returncode,
        }
        manifest["steps"][name] = {
            "status": "complete" if returncode == 0 else "retrying",
            "attempts": attempts,
            "command": current,
            "step_identity": step_identity,
            "output_state": file_state(output) if returncode == 0 else None,
        }
        write_json_atomic(manifest_path, manifest)
        if returncode == 0:
            if not output.is_file():
                raise RuntimeError(f"step {name} returned success without {output}")
            return
    manifest["steps"][name]["status"] = "failed"
    write_json_atomic(manifest_path, manifest)
    failures = sum(
        attempt.get("returncode") not in (None, 0) for attempt in attempts
    )
    raise RuntimeError(
        f"step {name} still has failed cases after {failures} failed attempts"
    )


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Run the reproducible full waypoint extraction and evaluation pipeline."
    )
    parser.add_argument("--run-dir", type=Path)
    parser.add_argument("--case-ids")
    parser.add_argument("--benchmark", type=Path, default=PROJECT_ROOT / "benchmark/benchmark.csv")
    parser.add_argument(
        "--cases-root",
        type=Path,
        default=Path.home() / "kernels/SyzPilot-experiments/cases",
    )
    parser.add_argument(
        "--configs-dir",
        type=Path,
        default=Path.home() / "kernels/SyzPilot-experiments/configs",
    )
    parser.add_argument(
        "--coverage-results", type=Path, default=DEFAULT_COVERAGE_RESULTS
    )
    parser.add_argument("--syzpilot-python", type=Path, default=DEFAULT_SYZPILOT_PYTHON)
    parser.add_argument("--agent-python", type=Path, default=DEFAULT_AGENT_PYTHON)
    parser.add_argument("--codex-home", type=Path, default=Path.home() / ".codex_jjy")
    parser.add_argument("--model", default="gpt-5.6-sol")
    parser.add_argument("--effort", choices=("low", "medium", "high", "xhigh"), default="xhigh")
    parser.add_argument("--max-concurrency", type=int, default=6)
    parser.add_argument("--agent-retries", type=int, default=2)
    parser.add_argument("--target-resolution-retries", type=int, default=2)
    parser.add_argument("--resolver-timeout-seconds", type=float, default=300.0)
    add_quality_weight_arguments(parser)
    args = parser.parse_args()
    if args.max_concurrency < 1:
        parser.error("--max-concurrency must be at least 1")
    if args.agent_retries < 0:
        parser.error("--agent-retries must be nonnegative")
    if args.target_resolution_retries < 0:
        parser.error("--target-resolution-retries must be nonnegative")
    if args.resolver_timeout_seconds <= 0:
        parser.error("--resolver-timeout-seconds must be positive")
    try:
        quality_weights = quality_weights_from_args(args)
    except ValueError as error:
        parser.error(str(error))
    for executable in (args.syzpilot_python, args.agent_python):
        if not executable.is_file():
            parser.error(f"Python executable not found: {executable}")
    if not args.coverage_results.is_file():
        parser.error(
            f"compatible KCOV results are required for a canonical run: "
            f"{args.coverage_results}"
        )

    git_sha, git_dirty = get_git_state()
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    run_dir = args.run_dir or (
        PROJECT_ROOT
        / "agent_analysis/waypoint_runs"
        / f"{timestamp}_{git_sha[:12]}"
    )
    cache_dir = run_dir / ".cache"
    logs_dir = run_dir / "logs"
    run_dir.mkdir(parents=True, exist_ok=True)
    cache_dir.mkdir(parents=True, exist_ok=True)
    logs_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = run_dir / "manifest.json"
    run_configuration = {
        "case_ids": args.case_ids or "all",
        "benchmark": str(args.benchmark.resolve()),
        "cases_root": str(args.cases_root.resolve()),
        "configs_dir": str(args.configs_dir.resolve()),
        "coverage_results": str(args.coverage_results.resolve()),
        "model": args.model,
        "effort": args.effort,
        "max_concurrency": args.max_concurrency,
        "codex_home": str(args.codex_home.expanduser().resolve()),
        "target_resolution_retries": args.target_resolution_retries,
        "resolver_timeout_seconds": args.resolver_timeout_seconds,
        "quality_weights": quality_weights.as_dict(),
    }
    if manifest_path.is_file():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        if manifest.get("configuration") != run_configuration:
            parser.error("existing run directory uses a different configuration")
    else:
        manifest = {
            "schema_version": "2.0",
            "run_id": run_dir.name,
            "created_utc": timestamp,
            "git_sha": git_sha,
            "git_dirty": git_dirty,
            "configuration": run_configuration,
            "steps": {},
        }
        write_json_atomic(manifest_path, manifest)

    paths = {
        "static_workbook_cache": cache_dir / "static_waypoints.xlsx",
        "static_json": run_dir / "static_waypoints.json",
        "agentic_extractions": run_dir / "agentic_extractions.jsonl",
        "scoring_input": run_dir / "blind_scoring_input.json",
        "blind_scores": run_dir / "blind_scores.jsonl",
        "evaluation_json": run_dir / "waypoints_evaluation.json",
        "evaluation_workbook": run_dir / "waypoints_evaluation.xlsx",
        "readme": run_dir / "README.md",
    }
    case_args = ["--case-ids", args.case_ids] if args.case_ids else []

    execute_step(
        "script_extraction",
        [
            str(args.syzpilot_python),
            "-m",
            "analyzer.evaluate_waypoints_extractor",
            "--benchmark",
            str(args.benchmark),
            "--cases-root",
            str(args.cases_root),
            "--configs-dir",
            str(args.configs_dir),
            "--output",
            str(paths["static_workbook_cache"]),
            "--json-output",
            str(paths["static_json"]),
            *case_args,
        ],
        [paths["static_workbook_cache"], paths["static_json"]],
        manifest,
        manifest_path,
        logs_dir / "script_extraction.log",
        input_paths=[
            args.benchmark,
            PROJECT_ROOT / "benchmark/waypoints_oracle.json",
            ANALYZER_DIR / "waypoints_extractor.py",
            ANALYZER_DIR / "get_targets.py",
            ANALYZER_DIR / "evaluate_waypoints_extractor.py",
        ],
        cleanup_paths=[
            paths["static_workbook_cache"].with_suffix(".partial.xlsx"),
            paths["static_workbook_cache"].with_suffix(".partial.xlsx").with_name(
                ".static_waypoints.partial.writing.xlsx"
            ),
            paths["static_json"].with_name(f".{paths['static_json'].name}.writing"),
        ],
        reuse_validator=lambda: static_artifact_matches_inputs(paths["static_json"]),
        allow_nonzero_with_outputs=True,
    )
    validate_static_artifact_for_pipeline(paths["static_json"])

    execute_agent_step(
        "agentic_extraction",
        [
            str(args.agent_python),
            "-m",
            "analyzer.agentic_waypoints_extractor",
            "--benchmark",
            str(args.benchmark),
            "--cases-root",
            str(args.cases_root),
            "--configs-dir",
            str(args.configs_dir),
            "--output",
            str(paths["agentic_extractions"]),
            "--codex-home",
            str(args.codex_home),
            "--resolver-python",
            str(args.syzpilot_python),
            "--model",
            args.model,
            "--effort",
            args.effort,
            "--max-concurrency",
            str(args.max_concurrency),
            "--target-resolution-retries",
            str(args.target_resolution_retries),
            "--resolver-timeout-seconds",
            str(args.resolver_timeout_seconds),
            *case_args,
        ],
        paths["agentic_extractions"],
        args.agent_retries,
        manifest,
        manifest_path,
        logs_dir / "agentic_extraction.log",
        input_paths=[
            args.benchmark,
            paths["static_json"],
            ANALYZER_DIR / "agentic_waypoints_extractor.py",
            ANALYZER_DIR / "agentic_target_validator.py",
            ANALYZER_DIR / "waypoint_schema.py",
            ANALYZER_DIR / "codex_blind_wrapper.sh",
            ANALYZER_DIR / "prompts/agentic_waypoint_extractor.md",
        ],
    )

    execute_step(
        "prepare_blind_scoring",
        [
            str(args.syzpilot_python),
            "-m",
            "analyzer.waypoint_evaluation",
            "prepare",
            "--static-json",
            str(paths["static_json"]),
            "--agentic-extractions",
            str(paths["agentic_extractions"]),
            "--output",
            str(paths["scoring_input"]),
        ],
        [paths["scoring_input"]],
        manifest,
        manifest_path,
        logs_dir / "prepare_blind_scoring.log",
        input_paths=[paths["static_json"], paths["agentic_extractions"]],
        reuse_validator=lambda: json_artifact_is_readable(paths["scoring_input"]),
    )

    execute_agent_step(
        "blind_scoring",
        [
            str(args.agent_python),
            "-m",
            "analyzer.agentic_waypoint_scorer",
            "--input",
            str(paths["scoring_input"]),
            "--output",
            str(paths["blind_scores"]),
            "--codex-home",
            str(args.codex_home),
            "--model",
            args.model,
            "--effort",
            args.effort,
            "--max-concurrency",
            str(args.max_concurrency),
            *case_args,
        ],
        paths["blind_scores"],
        args.agent_retries,
        manifest,
        manifest_path,
        logs_dir / "blind_scoring.log",
        input_paths=[
            paths["scoring_input"],
            ANALYZER_DIR / "agentic_waypoint_scorer.py",
            ANALYZER_DIR / "waypoint_schema.py",
            ANALYZER_DIR / "codex_blind_wrapper.sh",
            ANALYZER_DIR / "prompts/agentic_waypoint_scorer.md",
        ],
    )

    build_command = [
        str(args.syzpilot_python),
        "-m",
        "analyzer.waypoint_evaluation",
        "build",
        "--static-json",
        str(paths["static_json"]),
        "--static-workbook",
        str(paths["static_workbook_cache"]),
        "--agentic-extractions",
        str(paths["agentic_extractions"]),
        "--scoring-input",
        str(paths["scoring_input"]),
        "--blind-scores",
        str(paths["blind_scores"]),
        "--output",
        str(paths["evaluation_workbook"]),
        "--json-output",
        str(paths["evaluation_json"]),
        "--semantic-weight",
        str(quality_weights.semantic),
        "--hit-weight",
        str(quality_weights.hit_quality),
        "--length-weight",
        str(quality_weights.effective_length),
    ]
    build_command.extend(["--coverage-results", str(args.coverage_results)])
    execute_step(
        "build_evaluation",
        build_command,
        [paths["evaluation_json"], paths["evaluation_workbook"], paths["readme"]],
        manifest,
        manifest_path,
        logs_dir / "build_evaluation.log",
        input_paths=[
            paths["static_json"],
            paths["static_workbook_cache"],
            paths["agentic_extractions"],
            paths["scoring_input"],
            paths["blind_scores"],
            args.coverage_results,
            ANALYZER_DIR / "waypoint_evaluation.py",
            ANALYZER_DIR / "evaluate_waypoint_quality.py",
        ],
        reuse_validator=lambda: evaluation_artifacts_are_readable(
            paths["evaluation_json"],
            paths["evaluation_workbook"],
            paths["readme"],
        ),
    )

    manifest["status"] = "complete"
    manifest["artifacts"] = artifact_state(paths)
    manifest["coverage_reused"] = {
        "path": str(args.coverage_results.resolve()),
        "policy": "reuse supplied results unless the collector is rerun with --overwrite",
    }
    write_json_atomic(manifest_path, manifest)
    print(f"Canonical workbook: {paths['evaluation_workbook']}")
    print(f"Canonical JSON: {paths['evaluation_json']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
