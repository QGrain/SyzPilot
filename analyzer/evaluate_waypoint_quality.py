#!/usr/bin/env python3
"""Enrich a waypoint regression workbook with agentic and dynamic evidence."""

from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

if __package__:
    from .evaluate_waypoints_extractor import (
        ANALYZER_DIR,
        append_mapping_sheet,
        format_pc32,
        save_workbook_atomic,
    )
else:
    from evaluate_waypoints_extractor import (
        ANALYZER_DIR,
        append_mapping_sheet,
        format_pc32,
        save_workbook_atomic,
    )


if str(ANALYZER_DIR) not in sys.path:
    sys.path.insert(0, str(ANALYZER_DIR))

from get_targets import (  # noqa: E402
    action_fast_build,
    check_target,
    get_target_pc,
    normalize_target,
)


ZERO_PC64 = "0x0000000000000000"
ZERO_PC32 = "0x00000000"
AMD64_KCOV_CALL_INSTRUCTION_LEN = 5
COVERAGE_PC_SEMANTICS_ID = "syz-execprog-amd64-previous-instruction-pc-v1"
ALLOWED_EXTRACTION_TYPES = {
    "script extracted",
    "agentic extracted",
    "manual fixed",
    "missing input",
}


def json_list(values: Iterable[Any]) -> str:
    return json.dumps(list(values), separators=(",", ":"), ensure_ascii=True)


def load_json(path: Path) -> Any:
    with path.open(encoding="utf-8") as stream:
        return json.load(stream)


def load_reviews(paths: list[Path]) -> dict[str, dict[str, Any]]:
    reviews: dict[str, dict[str, Any]] = {}
    for path in paths:
        payload = load_json(path)
        if isinstance(payload, dict):
            entries = payload.get("reviews", payload.get("cases", payload))
        else:
            entries = payload
        if isinstance(entries, dict):
            entries = [dict(value, case_id=key) for key, value in entries.items()]
        if not isinstance(entries, list):
            raise ValueError(f"review file must contain a case list: {path}")
        for entry in entries:
            if not isinstance(entry, dict) or "case_id" not in entry:
                raise ValueError(f"invalid review entry in {path}: {entry!r}")
            case_id = str(entry["case_id"])
            if case_id in reviews:
                raise ValueError(f"duplicate review for case {case_id}")
            extraction_type = entry.get("extraction_type", "")
            if extraction_type not in ALLOWED_EXTRACTION_TYPES:
                raise ValueError(
                    f"invalid extraction_type for case {case_id}: {extraction_type!r}"
                )
            score = entry.get("agentic_quality_score")
            if score is not None and not 0 <= float(score) <= 100:
                raise ValueError(f"invalid agentic score for case {case_id}: {score}")
            reviews[case_id] = entry
    return reviews


def load_phase_inventory(path: Path) -> dict[str, dict[str, Any]]:
    payload = load_json(path)
    return {str(entry["case_id"]): entry for entry in payload.get("cases", [])}


def sheet_rows_by_id(sheet) -> dict[str, dict[str, Any]]:
    rows = sheet.iter_rows(values_only=True)
    headers = list(next(rows))
    return {
        str(row[headers.index("ID")]): dict(zip(headers, row))
        for row in rows
    }


def load_final_pc64(workbook) -> dict[str, list[str]]:
    sheet = workbook["Waypoints_Long"]
    rows = sheet.iter_rows(values_only=True)
    headers = list(next(rows))
    by_case: dict[str, list[tuple[int, str]]] = defaultdict(list)
    for values in rows:
        row = dict(zip(headers, values))
        if row.get("Stage") != "outlier_removal":
            continue
        by_case[str(row["Case ID"])].append(
            (int(row["Index target->entry"]), row.get("PC64") or ZERO_PC64)
        )
    return {
        case_id: [pc for _, pc in sorted(items)]
        for case_id, items in by_case.items()
    }


def parse_json_cell(value: Any) -> list[Any]:
    if value in (None, ""):
        return []
    parsed = json.loads(value)
    if not isinstance(parsed, list):
        raise ValueError(f"expected JSON list cell, got {type(parsed).__name__}")
    return parsed


def resolve_agentic_chain(
    kernel_dir: str, targets: list[str]
) -> tuple[list[str], list[str], list[str], list[str]]:
    if not targets:
        return [], [], [], []
    info = action_fast_build(
        kernel_dir, targets, rebuild=False, allow_partial=True
    )
    resolved_targets: list[str] = []
    pcs64: list[str] = []
    pcs32: list[str] = []
    errors: list[str] = []
    for target in targets:
        try:
            normalized = normalize_target(kernel_dir, target)
            resolved = check_target(info, normalized, kernel_dir, 0)
            pc64 = get_target_pc(info, resolved, kernel_dir)
            if pc64 is None:
                raise ValueError("resolved target has no PC")
            error = ""
        except Exception as exc:
            resolved = target
            pc64 = ZERO_PC64
            error = f"{target}: {type(exc).__name__}: {exc}"
        resolved_targets.append(resolved)
        pcs64.append(pc64)
        pcs32.append(format_pc32(pc64) or ZERO_PC32)
        errors.append(error)
    return resolved_targets, pcs64, pcs32, errors


def validate_evaluated_pcs(
    case_id: str,
    targets: list[str],
    pcs64: list[str],
    errors: list[str],
    allow_empty: bool = False,
) -> None:
    if not targets:
        if pcs64 or errors:
            raise ValueError(f"case {case_id} has misaligned empty evaluated chain")
        if not allow_empty:
            raise ValueError(f"case {case_id} has an empty evaluated chain")
        return
    if len(pcs64) != len(targets) or len(errors) != len(targets):
        raise ValueError(f"case {case_id} evaluated PC arrays are not aligned")
    resolution_errors = [error for error in errors if error]
    if resolution_errors:
        raise ValueError(
            f"case {case_id} has unresolved evaluated waypoints: {resolution_errors}"
        )
    pcs32 = [parse_pc64(pc) & 0xffffffff for pc in pcs64]
    zero_indices = [index for index, pc in enumerate(pcs32) if pc == 0]
    if zero_indices:
        raise ValueError(
            f"case {case_id} has zero evaluated PCs at indices {zero_indices}"
        )
    duplicates = sorted(pc for pc in set(pcs32) if pcs32.count(pc) > 1)
    if duplicates:
        formatted = [f"0x{pc:08x}" for pc in duplicates]
        raise ValueError(f"case {case_id} has duplicate evaluated PCs: {formatted}")


def load_coverage_results(
    path: Path | None,
) -> tuple[Path | None, dict[str, list[dict[str, Any]]], dict[str, Any]]:
    if path is None:
        return None, {}, {}
    by_case: dict[str, list[dict[str, Any]]] = defaultdict(list)
    with path.open(encoding="utf-8") as stream:
        for line in stream:
            if not line.strip():
                continue
            row = json.loads(line)
            by_case[str(row["case_id"])].append(row)
    for rows in by_case.values():
        rows.sort(key=lambda row: int(row.get("run_idx", 0)))
    manifest_path = path.parent / "manifest.json"
    manifest = load_json(manifest_path) if manifest_path.is_file() else {}
    return path.parent, dict(by_case), manifest


def read_pc_file(path: Path) -> set[int]:
    if not path.is_file():
        return set()
    pcs = set()
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        try:
            value = int(line.strip(), 16)
        except ValueError:
            continue
        if value:
            pcs.add(value)
    return pcs


def fuzzer_pc_to_coverfile_pc(pc: int) -> int:
    """Convert a KCOV return PC to syz-execprog's amd64 callsite PC."""
    if pc <= AMD64_KCOV_CALL_INSTRUCTION_LEN:
        return 0
    return pc - AMD64_KCOV_CALL_INSTRUCTION_LEN


def parse_pc64(pc: Any) -> int:
    try:
        return int(pc, 16)
    except (TypeError, ValueError):
        return 0


def dynamic_coverage_score(
    hit_count: int, waypoint_total: int, target_hit: bool
) -> float:
    """Score comparable PoC evidence while rewarding longer fully hit chains."""
    if waypoint_total < 1:
        raise ValueError("waypoint_total must be positive")
    if not 0 <= hit_count <= waypoint_total:
        raise ValueError("hit_count must be in [0, waypoint_total]")
    score = 5.0
    if hit_count:
        score += 10.0
    score += 35.0 * hit_count / waypoint_total
    score += 30.0 * hit_count / (hit_count + 3.0)
    score += 20.0 * int(target_hit)
    ceiling = 100.0 if target_hit else 75.0
    return round(min(score, ceiling), 1)


def coverage_compatibility_errors(
    _kernel_dir: str,
    runs: list[dict[str, Any]],
    manifest: dict[str, Any],
) -> list[str]:
    errors: list[str] = []
    if not runs:
        return errors
    if manifest.get("target_arch") != "amd64":
        errors.append(
            f"manifest target_arch={manifest.get('target_arch')!r}, expected 'amd64'"
        )
    if manifest.get("coverage_pc_semantics_id") != COVERAGE_PC_SEMANTICS_ID:
        errors.append("manifest coverage PC semantics do not match evaluator")
    for row in runs:
        if row.get("preflight_status") != "OK":
            continue
        run_label = f"run{row.get('run_idx', '?')}"
        if row.get("target_arch") != "amd64":
            errors.append(f"{run_label} target_arch mismatch")
        if row.get("coverage_pc_semantics_id") != COVERAGE_PC_SEMANTICS_ID:
            errors.append(f"{run_label} coverage PC semantics mismatch")
    return sorted(set(errors))


def coverage_union_path(
    coverage_root: Path,
    row: dict[str, Any],
    field: str,
    filename: str,
) -> Path | None:
    relative = row.get(field)
    if not relative:
        result_dir = row.get("result_dir", "")
        relative = f"{result_dir}/coverage/{filename}" if result_dir else ""
    return coverage_root / relative if relative else None


def coverage_artifact_errors(
    coverage_root: Path | None,
    runs: list[dict[str, Any]],
) -> list[str]:
    if not runs:
        return []
    if coverage_root is None:
        return ["coverage root is unavailable"]
    errors: list[str] = []
    artifacts = (
        ("calls_only_union_pc64", "calls_only.union.pc64", "calls_only_pc_count"),
        ("extra_union_pc64", "extra.union.pc64", "extra_pc_count"),
    )
    for row in runs:
        run_label = f"run{row.get('run_idx', '?')}"
        for path_field, filename, count_field in artifacts:
            path = coverage_union_path(coverage_root, row, path_field, filename)
            if path is None or not path.is_file():
                errors.append(f"{run_label} missing coverage artifact: {filename}")
                continue
            expected_count = row.get(count_field)
            actual_count = len(read_pc_file(path))
            if expected_count is None:
                errors.append(f"{run_label} missing {count_field}")
            elif int(expected_count) != actual_count:
                errors.append(
                    f"{run_label} {filename} count mismatch: "
                    f"metadata={expected_count}, file={actual_count}"
                )
    return sorted(set(errors))


def coverage_evidence(
    coverage_root: Path | None,
    runs: list[dict[str, Any]],
    pcs64: list[str],
    compatibility_errors: list[str] | None = None,
    configured_target_index: int | None = 0,
) -> dict[str, Any]:
    normalized_pcs = [parse_pc64(pc) for pc in pcs64]
    comparison_pcs = [fuzzer_pc_to_coverfile_pc(pc) for pc in normalized_pcs]
    if not runs:
        return {
            "PoC Evaluation Status": "NOT_RUN",
            "Coverage Run Count": 0,
            "Comparable Coverage Run Count": 0,
            "Coverage Statuses": "",
            "KASLR Statuses": "",
            "Target Reproduced": None,
            "Coverage Union PC Count": None,
            "Extra Coverage Union PC Count": None,
            "Waypoint Hit Count": None,
            "Waypoint Total": len(pcs64),
            "Waypoint Hit Ratio": None,
            "Target Hit": None,
            "Per-Waypoint Hits (target->entry)": "",
            "Per-Waypoint Extra-only Hits (target->entry)": "",
            "Coverage Comparison PCs64 (target->entry)": json_list(
                f"0x{pc:016x}" for pc in comparison_pcs
            ),
            "Dynamic Coverage Score": None,
            "Dynamic Evidence Confidence": "not available",
            "Missing Waypoint PCs": "",
            "Coverage Artifacts": "",
            "Extra Coverage Artifacts": "",
            "PoC Fidelity Statuses": "",
            "Artifact Compatibility Errors": "",
            "Hit Sources": {},
            "Extra Hit Sources": {},
        }

    compatibility_errors = list(compatibility_errors or [])
    preflight = {row.get("preflight_status", "") for row in runs}
    coverage_statuses = [row.get("coverage_status", "") for row in runs]
    kaslr_statuses = [row.get("kaslr_status", "unknown") for row in runs]
    fidelity_statuses = [row.get("poc_fidelity_status", "UNKNOWN") for row in runs]
    coverage_candidates = [
        row
        for row in runs
        if row.get("coverage_status") in {"COMPLETE", "PARTIAL", "EMPTY"}
    ]
    usable_runs = [
        row
        for row in coverage_candidates
        if row.get("kaslr_status") == "disabled"
    ]
    compatibility_errors.extend(
        coverage_artifact_errors(coverage_root, usable_runs)
    )
    compatibility_errors = sorted(set(compatibility_errors))
    comparable = bool(usable_runs) and not compatibility_errors
    coverage_union: set[int] = set()
    extra_union: set[int] = set()
    call_hit_sources: dict[int, list[str]] = defaultdict(list)
    extra_hit_sources: dict[int, list[str]] = defaultdict(list)
    artifacts = []
    extra_artifacts = []
    if coverage_root is not None:
        for row in usable_runs:
            result_dir = row.get("result_dir", "")
            path = coverage_union_path(
                coverage_root,
                row,
                "calls_only_union_pc64",
                "calls_only.union.pc64",
            )
            if path is not None:
                artifacts.append(str(path))
                coverage_union.update(read_pc_file(path))
            path = coverage_union_path(
                coverage_root,
                row,
                "extra_union_pc64",
                "extra.union.pc64",
            )
            if path is not None:
                extra_artifacts.append(str(path))
                extra_union.update(read_pc_file(path))
            raw_dir = coverage_root / result_dir / "coverage" / "raw"
            if raw_dir.is_dir():
                for raw_file in sorted(raw_dir.iterdir()):
                    raw_pcs = read_pc_file(raw_file)
                    source_map = (
                        extra_hit_sources
                        if raw_file.name.endswith(".extra")
                        else call_hit_sources
                    )
                    for pc in raw_pcs:
                        source_map[pc].append(
                            f"run{row.get('run_idx')}:{raw_file.name}"
                        )

    hits = [
        bool(pc and comparable and pc in coverage_union) for pc in comparison_pcs
    ]
    extra_only_hits = [
        bool(
            pc
            and comparable
            and pc not in coverage_union
            and pc in extra_union
        )
        for pc in comparison_pcs
    ]
    total = len(hits)
    hit_count = sum(hits)
    ratio = hit_count / total if total else 0.0
    target_hit = bool(
        configured_target_index is not None
        and 0 <= configured_target_index < len(hits)
        and hits[configured_target_index]
    )

    missing_poc = preflight == {"SKIPPED_MISSING_POC"}
    if missing_poc:
        evaluation_status = "SKIPPED_MISSING_POC"
    elif not coverage_candidates:
        evaluation_status = "NO_USABLE_COVERAGE_RUN"
    elif not usable_runs:
        evaluation_status = "BLOCKED_KASLR_NOT_DISABLED"
    elif compatibility_errors:
        evaluation_status = "BLOCKED_ARTIFACT_MISMATCH"
    elif not coverage_union and extra_union:
        evaluation_status = "EVALUATED_EXTRA_ONLY"
    elif not coverage_union:
        evaluation_status = "EVALUATED_NO_RECORDED_COVERAGE"
    else:
        evaluation_status = "EVALUATED"
    dynamic_score: float | None = None
    if evaluation_status.startswith("EVALUATED") and total:
        dynamic_score = dynamic_coverage_score(hit_count, total, target_hit)

    if not evaluation_status.startswith("EVALUATED"):
        confidence = "not available"
    elif not any(bool(row.get("target_reproduced")) for row in usable_runs):
        confidence = "non-reproducing lower-bound evidence"
    elif all(row.get("coverage_status") == "COMPLETE" for row in usable_runs) and all(
        row.get("poc_fidelity_status") == "EXACT" for row in usable_runs
    ):
        confidence = "complete-run evidence"
    elif any(
        row.get("poc_fidelity_status") == "DEGRADED_UNSUPPORTED_OPTIONS"
        for row in usable_runs
    ):
        confidence = "partial lower-bound evidence (PoC option fidelity degraded)"
    else:
        confidence = "partial lower-bound evidence"
    missing = [
        f"0x{pc:016x}" for pc, hit in zip(normalized_pcs, hits) if pc and not hit
    ]
    return {
        "PoC Evaluation Status": evaluation_status,
        "Coverage Run Count": len(runs),
        "Comparable Coverage Run Count": len(usable_runs),
        "Coverage Statuses": json_list(coverage_statuses),
        "KASLR Statuses": json_list(kaslr_statuses),
        "Target Reproduced": any(bool(row.get("target_reproduced")) for row in runs),
        "Coverage Union PC Count": len(coverage_union),
        "Extra Coverage Union PC Count": len(extra_union),
        "Waypoint Hit Count": hit_count if comparable else None,
        "Waypoint Total": total,
        "Waypoint Hit Ratio": round(ratio, 4) if comparable and total else None,
        "Target Hit": target_hit if comparable and total else None,
        "Per-Waypoint Hits (target->entry)": json_list(hits) if comparable else "",
        "Per-Waypoint Extra-only Hits (target->entry)": (
            json_list(extra_only_hits) if comparable else ""
        ),
        "Coverage Comparison PCs64 (target->entry)": json_list(
            f"0x{pc:016x}" for pc in comparison_pcs
        ),
        "Dynamic Coverage Score": dynamic_score,
        "Dynamic Evidence Confidence": confidence,
        "Missing Waypoint PCs": json_list(missing) if comparable else "",
        "Coverage Artifacts": "\n".join(artifacts),
        "Extra Coverage Artifacts": "\n".join(extra_artifacts),
        "PoC Fidelity Statuses": json_list(fidelity_statuses),
        "Artifact Compatibility Errors": "\n".join(compatibility_errors),
        "Hit Sources": {
            f"0x{fuzzer_pc:016x}": call_hit_sources.get(coverage_pc, [])
            for fuzzer_pc, coverage_pc in zip(normalized_pcs, comparison_pcs)
            if fuzzer_pc
        },
        "Extra Hit Sources": {
            f"0x{fuzzer_pc:016x}": extra_hit_sources.get(coverage_pc, [])
            for fuzzer_pc, coverage_pc in zip(normalized_pcs, comparison_pcs)
            if fuzzer_pc
        },
    }


def build_evaluation_rows(
    cases: dict[str, dict[str, Any]],
    script_pc64: dict[str, list[str]],
    reviews: dict[str, dict[str, Any]],
    inventory: dict[str, dict[str, Any]],
    coverage_root: Path | None,
    coverage_runs: dict[str, list[dict[str, Any]]],
    coverage_manifest: dict[str, Any],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    evaluation_rows = []
    waypoint_rows = []
    resolved_payload: dict[str, Any] = {
        "schema_version": "1.0",
        "score_semantics": {
            "agentic_quality_score": (
                "Subjective report-grounded review: target 30, section 20, causal "
                "coherence 25, observability 15, parsimony 10. Not a unique ground truth."
            ),
            "dynamic_coverage_score": (
                "5 + 10*I(h>0) + 35*h/n + 30*h/(h+3) + "
                "20*target_hit; capped at 75 on target miss and 100 otherwise. "
                "Extra coverage is reported separately because the fuzzer does not label it."
            ),
        },
        "cases": {},
    }
    for case_id in sorted(cases, key=int):
        case = cases[case_id]
        review = reviews.get(case_id)
        if review is None:
            raise ValueError(f"missing agentic review for case {case_id}")
        phase_info = inventory.get(case_id, {})
        extraction_type = review["extraction_type"]
        script_targets = (case.get("Final Waypoints (target->entry)") or "").splitlines()
        if extraction_type == "script extracted":
            targets = script_targets
            pcs64 = script_pc64.get(case_id, [])
            resolved_targets = (
                case.get("Final Resolved Targets (target->entry)") or ""
            ).splitlines()
            errors = ["" for _ in targets]
        else:
            targets = list(review.get("evaluated_waypoints_target_to_entry", []))
            resolved_targets, pcs64, _, errors = resolve_agentic_chain(
                case["Kernel Dir"], targets
            )
        pcs32 = [format_pc32(pc) or ZERO_PC32 for pc in pcs64]
        phases = list(review.get("causal_phases_target_to_entry", []))
        if len(phases) != len(targets):
            raise ValueError(
                f"case {case_id} has {len(targets)} targets but {len(phases)} phases"
            )
        while len(resolved_targets) < len(targets):
            resolved_targets.append(targets[len(resolved_targets)])
        while len(pcs64) < len(targets):
            pcs64.append(ZERO_PC64)
            pcs32.append(ZERO_PC32)
            errors.append("PC64 unavailable in script workbook")
        validate_evaluated_pcs(
            case_id,
            targets,
            pcs64,
            errors,
            allow_empty=extraction_type == "missing input",
        )

        case_runs = coverage_runs.get(case_id, [])
        compatibility_errors = coverage_compatibility_errors(
            case["Kernel Dir"],
            case_runs,
            coverage_manifest,
        )
        dynamic = coverage_evidence(
            coverage_root,
            case_runs,
            pcs64,
            compatibility_errors=compatibility_errors,
        )
        comparison_pcs64 = parse_json_cell(
            dynamic["Coverage Comparison PCs64 (target->entry)"]
        )
        hits = (
            json.loads(dynamic["Per-Waypoint Hits (target->entry)"])
            if dynamic["Per-Waypoint Hits (target->entry)"]
            else [None for _ in targets]
        )
        hit_sources = dynamic.pop("Hit Sources")
        extra_hit_sources = dynamic.pop("Extra Hit Sources")
        extra_hits = (
            json.loads(dynamic["Per-Waypoint Extra-only Hits (target->entry)"])
            if dynamic["Per-Waypoint Extra-only Hits (target->entry)"]
            else [None for _ in targets]
        )
        unresolved = [error for error in errors if error]
        row = {
            "ID": case_id,
            "Title": case.get("Title", ""),
            "Extraction Type": extraction_type,
            "Agentic Quality Score": review.get("agentic_quality_score", ""),
            "Agentic Confidence": review.get("confidence", ""),
            "Report Kind": phase_info.get("report_kind", ""),
            "Concurrency Class": phase_info.get("concurrency_class", ""),
            "Script Chain Verdict": review.get("script_chain_verdict", ""),
            "Target Verdict": review.get("target_verdict", ""),
            "Section Verdict": review.get("section_verdict", ""),
            "Required Causal Roles": json_list(review.get("required_roles", [])),
            "Observed Causal Roles": json_list(review.get("observed_roles", [])),
            "Evaluated Waypoints (target->entry)": "\n".join(targets),
            "Causal Phases (target->entry)": json_list(phases),
            "Evaluated Resolved Targets (target->entry)": json_list(resolved_targets),
            "Evaluated PCs64 (target->entry)": json_list(pcs64),
            "Evaluated PCs32 (target->entry)": json_list(pcs32),
            "Evaluated Fuzzer PCs32 (entry->target)": json_list(reversed(pcs32)),
            "Unresolved Evaluated Waypoints": "\n".join(unresolved),
            "Agentic Rationale": review.get("rationale", ""),
            "Unresolved Questions": "\n".join(
                review.get("unresolved_questions", [])
            ),
            **dynamic,
        }
        evaluation_rows.append(row)
        for index, (
            target,
            phase,
            resolved,
            pc64,
            pc32,
            comparison_pc64,
            error,
            hit,
            extra_hit,
        ) in enumerate(
            zip(
                targets,
                phases,
                resolved_targets,
                pcs64,
                pcs32,
                comparison_pcs64,
                errors,
                hits,
                extra_hits,
            )
        ):
            waypoint_rows.append(
                {
                    "Case ID": case_id,
                    "Extraction Type": extraction_type,
                    "Index target->entry": index,
                    "Index entry->target": len(targets) - index - 1,
                    "Causal Phase": phase,
                    "Evaluated Waypoint": target,
                    "Resolved Target": resolved,
                    "PC64": pc64,
                    "PC32": pc32,
                    "Coverfile Comparison PC64": comparison_pc64,
                    "Resolution Error": error,
                    "PoC Covered": "" if hit is None else hit,
                    "Coverage Sources": json_list(hit_sources.get(pc64, [])),
                    "PoC Extra-only Covered": (
                        "" if extra_hit is None else extra_hit
                    ),
                    "Extra Coverage Sources": json_list(
                        extra_hit_sources.get(pc64, [])
                    ),
                }
            )
        resolved_payload["cases"][case_id] = {
            "extraction_type": extraction_type,
            "agentic_quality_score": review.get("agentic_quality_score"),
            "evaluated_waypoints_target_to_entry": targets,
            "causal_phases_target_to_entry": phases,
            "resolved_targets_target_to_entry": resolved_targets,
            "pcs64_target_to_entry": pcs64,
            "pcs32_target_to_entry": pcs32,
            "resolution_errors": errors,
            "dynamic": dynamic,
        }
    return evaluation_rows, waypoint_rows, resolved_payload


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Add agentic and PoC KCOV evaluation sheets to a regression workbook."
    )
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--agentic-review", required=True, type=Path, nargs="+")
    parser.add_argument("--phase-inventory", required=True, type=Path)
    parser.add_argument("--coverage-results", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--resolved-json", type=Path)
    args = parser.parse_args()

    workbook = load_workbook(args.workbook)
    cases = sheet_rows_by_id(workbook["Cases"])
    script_pc64 = load_final_pc64(workbook)
    reviews = load_reviews(args.agentic_review)
    inventory = load_phase_inventory(args.phase_inventory)
    if set(reviews) != set(cases):
        missing = sorted(set(cases) - set(reviews), key=int)
        extra = sorted(set(reviews) - set(cases), key=int)
        raise ValueError(f"review/workbook case mismatch: missing={missing}, extra={extra}")
    coverage_root, coverage_runs, coverage_manifest = load_coverage_results(
        args.coverage_results
    )
    evaluation_rows, waypoint_rows, resolved = build_evaluation_rows(
        cases,
        script_pc64,
        reviews,
        inventory,
        coverage_root,
        coverage_runs,
        coverage_manifest,
    )
    for sheet_name in ("Evaluation", "Evaluation_Waypoints"):
        if sheet_name in workbook.sheetnames:
            workbook.remove(workbook[sheet_name])
    append_mapping_sheet(workbook, "Evaluation", evaluation_rows)
    append_mapping_sheet(workbook, "Evaluation_Waypoints", waypoint_rows)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    save_workbook_atomic(workbook, args.output)
    resolved_path = args.resolved_json or args.output.with_suffix(".evaluation.json")
    resolved_path.write_text(
        json.dumps(resolved, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    print(f"Wrote {args.output}")
    print(f"Wrote {resolved_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
