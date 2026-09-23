#!/usr/bin/env python3
"""Run the waypoint extractor as a reproducible, phase-level XLSX regression."""

from __future__ import annotations

import argparse
import csv
import json
import os
import platform
import re
import subprocess
import sys
import traceback
import uuid
from collections import Counter
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from statistics import mean, median
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
ANALYZER_DIR = PROJECT_ROOT / "analyzer"
if str(ANALYZER_DIR) not in sys.path:
    sys.path.insert(0, str(ANALYZER_DIR))

from get_targets import (  # noqa: E402
    action_fast_build,
    check_target,
    get_vmlinux_cache_identity,
    get_target_pc,
    normalize_target,
    resolve_function_for_location,
)
from waypoints_extractor import WaypointsExtractor  # noqa: E402


SCHEMA_VERSION = "1.0"
CANONICAL_SCHEMA_VERSION = "2.0"
CHAIN_DIRECTION = "target_to_entry"
ZERO_PC64 = "0x0000000000000000"
ZERO_PC32 = "0x00000000"


STAGES = (
    ("call_trace", "After Call Trace Extraction", "paper"),
    ("trace_sanitization", "After Trace Sanitization", "paper"),
    ("bb_resolution", "After Instrumentation and BB Resolution", "implementation"),
    ("complexity_filtering", "After Complexity Filtering", "paper"),
    ("trivial_stage_merging", "After Trivial Stage Merging", "paper"),
    ("greedy_pruning", "After Greedy Pruning", "paper"),
    ("outlier_removal", "After Outlier Removal", "engineering"),
)
STAGE_LABELS = {key: label for key, label, _ in STAGES}


@dataclass(frozen=True)
class NodeSnapshot:
    func_name: str
    location: str
    is_inline: bool
    original_index: int
    bb_offset: int | None
    bb_count: int | None
    hot_entry_score: float | None
    value_score: float | None

    @property
    def target(self) -> str:
        return f"{self.func_name}@{self.location}"

    @property
    def identity(self) -> tuple[str, str, int]:
        return self.func_name, self.location, self.original_index


@dataclass
class StageSnapshot:
    key: str
    nodes: list[NodeSnapshot]
    resolved_targets: list[str] = field(default_factory=list)
    pcs64: list[str | None] = field(default_factory=list)
    pcs32: list[str | None] = field(default_factory=list)
    resolution_errors: list[str] = field(default_factory=list)

    @property
    def targets(self) -> list[str]:
        return [node.target for node in self.nodes]


@dataclass
class ValidationResult:
    check: str
    severity: str
    passed: bool
    details: str = ""


@dataclass
class CaseResult:
    metadata: dict[str, str]
    report_type: str = ""
    report_sections: list[str] = field(default_factory=list)
    title_bug_func: str = ""
    bug_position_func: str = ""
    bug_position_source_target: str = ""
    bug_position_instrumentation_target: str = ""
    bug_position_pc64: str | None = None
    bug_position_pc32: str | None = None
    bug_position_resolution_class: str = "unavailable"
    status: str = "pending"
    failed_stage: str = ""
    error: str = ""
    warnings: list[str] = field(default_factory=list)
    snapshots: dict[str, StageSnapshot] = field(default_factory=dict)
    validations: list[ValidationResult] = field(default_factory=list)
    bug_position_exact_relation: str = "unavailable"
    bug_position_exact_index: int | None = None
    bug_position_function_relation: str = "unavailable"
    bug_position_function_index: int | None = None


def file_state(path: Path) -> dict[str, Any] | None:
    """Return a cheap identity for local resume checks without reading the file."""
    if not path.is_file():
        return None
    stat = path.stat()
    return {
        "path": str(path.resolve()),
        "device": stat.st_dev,
        "inode": stat.st_ino,
        "size": stat.st_size,
        "mtime_ns": stat.st_mtime_ns,
        "ctime_ns": stat.st_ctime_ns,
    }


def get_git_state() -> tuple[str, bool]:
    try:
        sha = subprocess.check_output(
            ["git", "rev-parse", "HEAD"], cwd=PROJECT_ROOT, text=True
        ).strip()
        dirty = bool(
            subprocess.check_output(
                ["git", "status", "--porcelain"], cwd=PROJECT_ROOT, text=True
            ).strip()
        )
        return sha, dirty
    except (OSError, subprocess.CalledProcessError):
        return "unknown", True


def normalize_function_name(name: str) -> str:
    """Normalize compiler-generated suffixes for audit metrics only."""
    suffix_pattern = r"(?:\.(?:isra|constprop|part)\.\d+|\.cold(?:\.\d+)?)$"
    previous = None
    while name != previous:
        previous = name
        name = re.sub(suffix_pattern, "", name)
    return name


def function_name_from_resolved_target(target: str) -> str:
    return target.split("@", 1)[0].strip()


def format_pc32(pc64: str | None) -> str | None:
    if pc64 is None:
        return None
    try:
        return f"0x{int(pc64, 16) & 0xFFFFFFFF:08x}"
    except (TypeError, ValueError):
        return ZERO_PC32


def ordered_subsequence(child: Iterable[Any], parent: Iterable[Any]) -> bool:
    parent_iter = iter(parent)
    return all(any(candidate == item for candidate in parent_iter) for item in child)


def capture_snapshot(extractor: WaypointsExtractor, key: str) -> StageSnapshot:
    """Copy the current linked list in the workbook's target-to-entry order."""
    nodes = [
        NodeSnapshot(
            func_name=node.func_name,
            location=node.location,
            is_inline=node.is_inline,
            original_index=node.original_index,
            bb_offset=node.bb_offset,
            bb_count=node.bb_count,
            hot_entry_score=node.hot_entry_score,
            value_score=node.value_score,
        )
        for node in extractor.waypoints
    ]
    nodes.reverse()
    return StageSnapshot(key=key, nodes=nodes)


def resolve_snapshot(
    snapshot: StageSnapshot,
    kernel_dir: Path,
    instrumentation_info: dict[str, Any] | None,
    core_scope: set[tuple[str, str, int]],
) -> None:
    for node in snapshot.nodes:
        target = node.target
        resolved_target = target
        pc64: str | None = ZERO_PC64
        error = ""
        if node.identity not in core_scope:
            snapshot.resolved_targets.append(resolved_target)
            snapshot.pcs64.append(None)
            snapshot.pcs32.append(None)
            snapshot.resolution_errors.append(
                f"{target}: not attempted because the node was removed before the "
                "core instrumentation build"
            )
            continue
        try:
            if instrumentation_info is None:
                raise ValueError("instrumentation information is unavailable")
            normalized_target = normalize_target(str(kernel_dir), target)
            resolved_target = check_target(
                instrumentation_info, normalized_target, str(kernel_dir), 0
            )
            resolved_pc = get_target_pc(
                instrumentation_info, resolved_target, str(kernel_dir)
            )
            if resolved_pc is None:
                raise ValueError(f"cannot resolve a PC for {resolved_target}")
            pc64 = resolved_pc
        except Exception as exc:  # A failed node must not discard the case row.
            error = f"{target}: {type(exc).__name__}: {exc}"
        snapshot.resolved_targets.append(resolved_target)
        snapshot.pcs64.append(pc64)
        snapshot.pcs32.append(format_pc32(pc64))
        snapshot.resolution_errors.append(error)


def classify_location_relation(
    snapshot: StageSnapshot | None, location: str
) -> tuple[str, int | None]:
    if snapshot is None or not location:
        return "unavailable", None
    normalized_location = os.path.normpath(location)
    for index, node in enumerate(snapshot.nodes):
        if os.path.normpath(node.location) == normalized_location:
            return ("target" if index == 0 else "toward_entry"), index
    return "absent", None


def classify_function_relation(
    snapshot: StageSnapshot | None, func_name: str
) -> tuple[str, int | None]:
    if snapshot is None or not func_name:
        return "unavailable", None
    normalized_func = normalize_function_name(func_name)
    for index, node in enumerate(snapshot.nodes):
        if normalize_function_name(node.func_name) == normalized_func:
            return ("target" if index == 0 else "toward_entry"), index
    return "absent", None


def load_benchmark(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as stream:
        return list(csv.DictReader(stream))


def detect_report_sections(report_path: Path) -> list[str]:
    text = report_path.read_text(encoding="utf-8", errors="replace")
    markers = (
        "Call Trace:",
        "Allocated by task",
        "Freed by task",
        "The buggy address belongs to the object at",
        "The buggy address is located",
        "Last potentially related work creation:",
        "Second to last potentially related work creation:",
        "page last allocated via order",
        "page last free stack trace:",
    )
    return [marker for marker in markers if marker in text]


def load_oracle(path: Path | None) -> dict[str, dict[str, Any]]:
    if path is None:
        return {}
    with path.open(encoding="utf-8") as stream:
        payload = json.load(stream)
    if payload.get("schema_version") != SCHEMA_VERSION:
        raise ValueError(
            f"unsupported oracle schema: {payload.get('schema_version')!r}"
        )
    cases = payload.get("cases")
    if not isinstance(cases, dict):
        raise ValueError("oracle 'cases' must be an object keyed by benchmark ID")
    list_fields = {
        "allowed_fallback_targets",
        "required_nodes",
        "allowed_nodes",
        "forbidden_nodes",
        "forbidden_report_sections",
    }
    string_fields = {
        "expected_target",
        "target_verdict",
        "section_verdict",
        "chain_verdict",
        "reviewer",
        "review_date",
        "notes",
        "accepted_exception",
    }
    validated = {}
    for case_id, entry in cases.items():
        if not isinstance(entry, dict):
            raise ValueError(f"oracle case {case_id!r} must be an object")
        for field_name in list_fields:
            value = entry.get(field_name, [])
            if not isinstance(value, list) or not all(
                isinstance(item, str) for item in value
            ):
                raise ValueError(
                    f"oracle case {case_id!r} field {field_name!r} "
                    "must be an array of strings"
                )
        for field_name in string_fields:
            if not isinstance(entry.get(field_name, ""), str):
                raise ValueError(
                    f"oracle case {case_id!r} field {field_name!r} must be a string"
                )
        validated[str(case_id)] = entry
    return validated


def record_snapshot(
    result: CaseResult, extractor: WaypointsExtractor, stage_key: str
) -> None:
    result.snapshots[stage_key] = capture_snapshot(extractor, stage_key)


def run_case(
    row: dict[str, str],
    cases_root: Path,
    configs_dir: Path,
    oracle_entry: dict[str, Any] | None = None,
) -> CaseResult:
    case_id = row["ID"]
    kernel_dir = cases_root / f"case_{case_id}"
    title_path = configs_dir / f"case_{case_id}.title"
    report_path = configs_dir / f"case_{case_id}.report"
    metadata = dict(row)
    metadata.update(
        {
            "Kernel Dir": str(kernel_dir),
            "Title Path": str(title_path),
            "Report Path": str(report_path),
        }
    )
    result = CaseResult(metadata=metadata)
    if report_path.is_file():
        result.report_sections = detect_report_sections(report_path)

    missing_inputs = [
        name
        for name, path in (
            ("kernel_dir", kernel_dir),
            ("title", title_path),
            ("report", report_path),
        )
        if not path.exists()
    ]
    if missing_inputs:
        result.status = "missing_input"
        result.failed_stage = "input_validation"
        result.error = f"Missing inputs: {', '.join(missing_inputs)}"
        validate_case(result, oracle_entry)
        return result

    extractor: WaypointsExtractor | None = None
    instrumentation_info: dict[str, Any] | None = None
    current_stage = "call_trace"
    try:
        extractor = WaypointsExtractor(
            str(kernel_dir), str(title_path), str(report_path), debug=False
        )
        result.report_type = extractor.report_type
        result.title_bug_func = extractor.bug_func
        record_snapshot(result, extractor, "call_trace")

        current_stage = "trace_sanitization"
        extractor.trace_sanitization()
        record_snapshot(result, extractor, current_stage)

        current_stage = "bb_resolution"
        target_to_entry = extractor.waypoints.recover_calltrace().splitlines()
        instrumentation_info = action_fast_build(
            str(kernel_dir), target_to_entry, False, allow_partial=True
        )
        extractor.load_bb_info(instrumentation_info)
        record_snapshot(result, extractor, current_stage)

        current_stage = "complexity_filtering"
        extractor.complexity_based_filtering()
        record_snapshot(result, extractor, current_stage)

        current_stage = "trivial_stage_merging"
        extractor.non_trivial_stage_separation()
        record_snapshot(result, extractor, current_stage)

        extractor.waypoints.calc_waypoints_value_scores(
            extractor.VALUE_SCORE_ALPHA,
            extractor.VALUE_SCORE_BETA,
            extractor.VALUE_SCORE_GAMMA,
        )
        current_stage = "greedy_pruning"
        extractor.recursively_greedy_triming()
        record_snapshot(result, extractor, current_stage)

        current_stage = "outlier_removal"
        extractor.remove_outliers()
        record_snapshot(result, extractor, current_stage)
        result.status = "ok"
    except (Exception, SystemExit) as exc:
        result.status = "error"
        result.failed_stage = current_stage
        result.error = (
            f"{type(exc).__name__}: {exc}\n"
            f"{traceback.format_exc(limit=8)}"
        )

    sanitized_snapshot = result.snapshots.get("trace_sanitization")
    core_scope = {
        node.identity for node in sanitized_snapshot.nodes
    } if sanitized_snapshot else set()
    for snapshot in result.snapshots.values():
        resolve_snapshot(snapshot, kernel_dir, instrumentation_info, core_scope)

    try:
        result.bug_position_source_target = resolve_function_for_location(
            str(kernel_dir), row.get("Bug Position", "")
        )
        result.bug_position_func = function_name_from_resolved_target(
            result.bug_position_source_target
        )
    except Exception as exc:
        result.warnings.append(
            f"Bug Position source resolution: {type(exc).__name__}: {exc}"
        )

    if result.bug_position_source_target:
        try:
            if instrumentation_info is None:
                raise ValueError("instrumentation information is unavailable")
            normalized_bug_target = normalize_target(
                str(kernel_dir), result.bug_position_source_target
            )
            result.bug_position_instrumentation_target = check_target(
                instrumentation_info, normalized_bug_target, str(kernel_dir), 0
            )
            result.bug_position_pc64 = get_target_pc(
                instrumentation_info,
                result.bug_position_instrumentation_target,
                str(kernel_dir),
            )
            if result.bug_position_pc64 is None:
                raise ValueError("resolved instrumentation target has no PC")
            result.bug_position_pc32 = format_pc32(result.bug_position_pc64)
            if result.bug_position_instrumentation_target == normalized_bug_target:
                result.bug_position_resolution_class = "exact_instrumentation_site"
            elif function_name_from_resolved_target(
                result.bug_position_instrumentation_target
            ) == function_name_from_resolved_target(normalized_bug_target):
                result.bug_position_resolution_class = "previous_instrumentation_site"
            else:
                result.bug_position_resolution_class = "caller_fallback"
        except Exception as exc:
            result.bug_position_instrumentation_target = ""
            result.bug_position_pc64 = ZERO_PC64
            result.bug_position_pc32 = ZERO_PC32
            result.bug_position_resolution_class = "unresolved"
            result.warnings.append(
                f"Bug Position instrumentation resolution: {type(exc).__name__}: {exc}"
            )

    final_snapshot = result.snapshots.get("outlier_removal")
    result.bug_position_exact_relation, result.bug_position_exact_index = (
        classify_location_relation(final_snapshot, row.get("Bug Position", ""))
    )
    (
        result.bug_position_function_relation,
        result.bug_position_function_index,
    ) = classify_function_relation(final_snapshot, result.bug_position_func)
    validate_case(result, oracle_entry)
    return result


def validate_case(
    result: CaseResult, oracle_entry: dict[str, Any] | None = None
) -> None:
    def add(check: str, severity: str, passed: bool, details: str = "") -> None:
        result.validations.append(
            ValidationResult(check, severity, bool(passed), details)
        )

    add(
        "pipeline_completed",
        "error",
        result.status == "ok",
        result.error or result.status,
    )

    snapshots = [result.snapshots.get(key) for key, _, _ in STAGES]
    present_snapshots = [snapshot for snapshot in snapshots if snapshot is not None]
    deletion_only = all(
        ordered_subsequence(
            [node.identity for node in child.nodes],
            [node.identity for node in parent.nodes],
        )
        for parent, child in zip(present_snapshots, present_snapshots[1:])
    )
    add(
        "phase_outputs_are_ordered_subsequences",
        "error",
        deletion_only,
        "Each filtering phase must only delete nodes and preserve relative order.",
    )

    final_snapshot = result.snapshots.get("outlier_removal")
    if final_snapshot is None:
        add("final_chain_available", "error", False, "Final stage was not reached.")
        return

    add("final_chain_available", "error", True)
    final_count = len(final_snapshot.nodes)
    add(
        "final_chain_nonempty",
        "error",
        final_count > 0,
        f"count={final_count}",
    )
    add(
        "final_count_at_most_kmax",
        "error",
        final_count <= WaypointsExtractor.MAX_WAYPOINTS_LENGTH,
        f"count={final_count}, K_max={WaypointsExtractor.MAX_WAYPOINTS_LENGTH}",
    )
    bb_resolved = result.snapshots.get("bb_resolution")
    min_expected = min(
        len(bb_resolved.nodes) if bb_resolved else 0,
        WaypointsExtractor.MIN_WAYPOINTS_LENGTH,
    )
    add(
        "final_count_preserves_available_kmin",
        "warning",
        final_count >= min_expected,
        f"count={final_count}, expected_at_least={min_expected}",
    )
    add(
        "final_waypoint_pc_lengths_align",
        "error",
        final_count
        == len(final_snapshot.resolved_targets)
        == len(final_snapshot.pcs64)
        == len(final_snapshot.pcs32),
    )
    invalid_frames = [
        node.target
        for node in final_snapshot.nodes
        if "@" not in node.target or not re.search(r":\d+$", node.location)
    ]
    add(
        "final_frames_are_structured_locations",
        "error",
        not invalid_frames,
        "; ".join(invalid_frames),
    )
    unresolved_pc_targets = [
        node.target
        for node, pc in zip(final_snapshot.nodes, final_snapshot.pcs64)
        if pc in (None, ZERO_PC64)
    ]
    add(
        "final_pcs_resolve_nonzero",
        "error",
        not unresolved_pc_targets,
        "; ".join(unresolved_pc_targets),
    )
    nonzero_pc32 = [
        pc for pc in final_snapshot.pcs32 if pc not in (None, ZERO_PC32)
    ]
    duplicates = sorted(pc for pc, count in Counter(nonzero_pc32).items() if count > 1)
    add(
        "final_pc32_values_are_unique",
        "error",
        not duplicates,
        json.dumps(duplicates),
    )
    title_targeted = bool(final_snapshot.nodes) and (
        normalize_function_name(final_snapshot.nodes[0].func_name)
        == normalize_function_name(result.title_bug_func)
    )
    add(
        "title_bug_function_is_terminal",
        "warning",
        title_targeted,
        result.title_bug_func,
    )
    add(
        "bug_position_function_is_terminal",
        "warning",
        result.bug_position_function_relation == "target",
        result.bug_position_function_relation,
    )
    if result.bug_position_pc32 not in (None, ZERO_PC32) and final_snapshot.pcs32:
        add(
            "bug_position_pc_matches_terminal_pc",
            "warning",
            result.bug_position_pc32 == final_snapshot.pcs32[0],
            (
                f"bug_position_pc={result.bug_position_pc32}; "
                f"terminal_pc={final_snapshot.pcs32[0]}"
            ),
        )
    if oracle_entry:
        expected_target = str(oracle_entry.get("expected_target", "")).strip()
        allowed_fallbacks = {
            str(value).strip()
            for value in oracle_entry.get("allowed_fallback_targets", [])
            if str(value).strip()
        }
        if expected_target:
            actual_targets = set()
            if final_snapshot.nodes:
                actual_targets.add(final_snapshot.nodes[0].target)
            if final_snapshot.resolved_targets:
                actual_targets.add(final_snapshot.resolved_targets[0])
            add(
                "oracle_terminal_matches",
                "error",
                bool(actual_targets & ({expected_target} | allowed_fallbacks)),
                f"expected={expected_target}; actual={sorted(actual_targets)}",
            )
        final_targets = set(final_snapshot.targets)
        required_nodes = {
            str(value).strip()
            for value in oracle_entry.get("required_nodes", [])
            if str(value).strip()
        }
        forbidden_nodes = {
            str(value).strip()
            for value in oracle_entry.get("forbidden_nodes", [])
            if str(value).strip()
        }
        if required_nodes:
            missing_required = sorted(required_nodes - final_targets)
            add(
                "oracle_required_nodes_present",
                "error",
                not missing_required,
                json.dumps(missing_required),
            )
        if forbidden_nodes:
            present_forbidden = sorted(forbidden_nodes & final_targets)
            add(
                "oracle_forbidden_nodes_absent",
                "error",
                not present_forbidden,
                json.dumps(present_forbidden),
            )


def json_list(values: Iterable[Any]) -> str:
    return json.dumps(list(values), ensure_ascii=True, separators=(",", ":"))


def chain_text(snapshot: StageSnapshot | None) -> str:
    return "\n".join(snapshot.targets) if snapshot else ""


def operational_final_snapshot(snapshot: StageSnapshot | None) -> StageSnapshot | None:
    """Apply the CLI's zero-PC and global PC32 deduplication to the final chain."""
    if snapshot is None:
        return None
    retained_indices = []
    seen_pc32: set[int] = set()
    for index, pc in enumerate(snapshot.pcs64):
        try:
            pc32 = int(pc, 16) & 0xFFFFFFFF
        except (TypeError, ValueError):
            pc32 = 0
        if not pc32 or pc32 in seen_pc32:
            continue
        seen_pc32.add(pc32)
        retained_indices.append(index)
    return StageSnapshot(
        key="operational_final",
        nodes=[snapshot.nodes[index] for index in retained_indices],
        resolved_targets=[snapshot.resolved_targets[index] for index in retained_indices],
        pcs64=[snapshot.pcs64[index] for index in retained_indices],
        pcs32=[snapshot.pcs32[index] for index in retained_indices],
        resolution_errors=[
            snapshot.resolution_errors[index] for index in retained_indices
        ],
    )


def stage_case_columns(stage_label: str) -> list[str]:
    return [
        f"{stage_label} Waypoints (target->entry)",
        f"{stage_label} Listed PCs (target->entry)",
        f"{stage_label} Resolved Targets (target->entry)",
        f"{stage_label} Count",
        f"{stage_label} PC Resolution Errors",
    ]


BASE_CASE_COLUMNS = [
    "ID",
    "SyzDirect ID",
    "Syzbot Bug ID",
    "Title",
    "Version",
    "Commit",
    "Bug Position",
    "Bug Position Function",
    "Bug Position Source Target",
    "Bug Position Instrumentation Target",
    "Bug Position PC64",
    "Bug Position PC32",
    "Bug Position Resolution Class",
    "Title Bug Function",
    "Report Type",
    "Detected Report Sections",
    "Status",
    "Failed Stage",
    "Error",
    "Warnings",
    "Kernel Dir",
    "Report Path",
]
FINAL_CASE_COLUMNS = [
    "Final Waypoints (target->entry)",
    "Final Listed PCs (target->entry)",
    "Final Resolved Targets (target->entry)",
    "Fuzzer target_pcs (entry->target)",
    "Final Count",
    "Final Terminal Frame",
    "Final Terminal PC64",
    "Final Terminal PC32",
    "Bug Position Exact Relation",
    "Bug Position Exact Index (target=0)",
    "Bug Position Function Relation",
    "Bug Position Function Index (target=0)",
    "Validation Error Count",
    "Validation Warning Count",
    "Validation Failure Signature",
]


def case_row(result: CaseResult) -> dict[str, Any]:
    row: dict[str, Any] = {
        "ID": result.metadata.get("ID", ""),
        "SyzDirect ID": result.metadata.get("SyzDirect_ID", ""),
        "Syzbot Bug ID": result.metadata.get("Syzbot Bug ID", ""),
        "Title": result.metadata.get("Title", ""),
        "Version": result.metadata.get("Version", ""),
        "Commit": result.metadata.get("Commit", ""),
        "Bug Position": result.metadata.get("Bug Position", ""),
        "Bug Position Function": result.bug_position_func,
        "Bug Position Source Target": result.bug_position_source_target,
        "Bug Position Instrumentation Target": result.bug_position_instrumentation_target,
        "Bug Position PC64": result.bug_position_pc64,
        "Bug Position PC32": result.bug_position_pc32,
        "Bug Position Resolution Class": result.bug_position_resolution_class,
        "Title Bug Function": result.title_bug_func,
        "Report Type": result.report_type,
        "Detected Report Sections": "\n".join(result.report_sections),
        "Status": result.status,
        "Failed Stage": result.failed_stage,
        "Error": result.error,
        "Warnings": "\n".join(result.warnings),
        "Kernel Dir": result.metadata.get("Kernel Dir", ""),
        "Report Path": result.metadata.get("Report Path", ""),
    }
    for key, label, _ in STAGES:
        snapshot = result.snapshots.get(key)
        row.update(
            {
                f"{label} Waypoints (target->entry)": chain_text(snapshot),
                f"{label} Listed PCs (target->entry)": (
                    json_list(snapshot.pcs32) if snapshot else ""
                ),
                f"{label} Resolved Targets (target->entry)": (
                    "\n".join(snapshot.resolved_targets) if snapshot else ""
                ),
                f"{label} Count": len(snapshot.nodes) if snapshot else "",
                f"{label} PC Resolution Errors": (
                    "\n".join(error for error in snapshot.resolution_errors if error)
                    if snapshot
                    else ""
                ),
            }
        )

    final_snapshot = operational_final_snapshot(
        result.snapshots.get("outlier_removal")
    )
    final_nodes = final_snapshot.nodes if final_snapshot else []
    error_count = sum(
        not validation.passed and validation.severity == "error"
        for validation in result.validations
    )
    warning_count = sum(
        not validation.passed and validation.severity == "warning"
        for validation in result.validations
    )
    row.update(
        {
            "Final Waypoints (target->entry)": chain_text(final_snapshot),
            "Final Listed PCs (target->entry)": (
                json_list(final_snapshot.pcs32) if final_snapshot else ""
            ),
            "Final Resolved Targets (target->entry)": (
                "\n".join(final_snapshot.resolved_targets)
                if final_snapshot
                else ""
            ),
            "Fuzzer target_pcs (entry->target)": (
                json_list(reversed(final_snapshot.pcs32)) if final_snapshot else ""
            ),
            "Final Count": len(final_nodes) if final_snapshot else "",
            "Final Terminal Frame": final_nodes[0].target if final_nodes else "",
            "Final Terminal PC64": (
                final_snapshot.pcs64[0] if final_snapshot and final_snapshot.pcs64 else ""
            ),
            "Final Terminal PC32": (
                final_snapshot.pcs32[0] if final_snapshot and final_snapshot.pcs32 else ""
            ),
            "Bug Position Exact Relation": result.bug_position_exact_relation,
            "Bug Position Exact Index (target=0)": result.bug_position_exact_index,
            "Bug Position Function Relation": result.bug_position_function_relation,
            "Bug Position Function Index (target=0)": result.bug_position_function_index,
            "Validation Error Count": error_count,
            "Validation Warning Count": warning_count,
            "Validation Failure Signature": json_list(
                f"{validation.severity}:{validation.check}"
                for validation in result.validations
                if not validation.passed
            ),
        }
    )
    return row


def style_sheet(sheet, freeze_panes: str = "A2") -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = freeze_panes
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        letter = get_column_letter(column_cells[0].column)
        max_length = max(
            len(str(cell.value).split("\n", 1)[0]) if cell.value is not None else 0
            for cell in column_cells
        )
        sheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 45)
        for cell in column_cells[1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def append_mapping_sheet(workbook: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(title)
    headers = list(rows[0].keys()) if rows else []
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    style_sheet(sheet)


def build_summary_rows(results: list[CaseResult]) -> list[dict[str, Any]]:
    status_counts = Counter(result.status for result in results)
    exact_counts = Counter(result.bug_position_exact_relation for result in results)
    function_counts = Counter(result.bug_position_function_relation for result in results)
    final_counts = [
        len(result.snapshots["outlier_removal"].nodes)
        for result in results
        if "outlier_removal" in result.snapshots
    ]
    rows: list[dict[str, Any]] = []

    def add(category: str, metric: str, value: Any, note: str = "") -> None:
        rows.append(
            {"Category": category, "Metric": metric, "Value": value, "Note": note}
        )

    add("Coverage", "Benchmark rows", len(results))
    for status, count in sorted(status_counts.items()):
        add("Coverage", f"Status: {status}", count)
    for relation, count in sorted(exact_counts.items()):
        add("Bug Position", f"Exact relation: {relation}", count)
    for relation, count in sorted(function_counts.items()):
        add("Bug Position", f"Function relation: {relation}", count)
    if final_counts:
        add("Final length", "Minimum", min(final_counts))
        add("Final length", "Maximum", max(final_counts))
        add("Final length", "Mean", round(mean(final_counts), 3))
        add("Final length", "Median", median(final_counts))
    for key, label, _ in STAGES:
        counts = [
            len(result.snapshots[key].nodes)
            for result in results
            if key in result.snapshots
        ]
        if counts:
            add("Stage length", label, round(mean(counts), 3), f"n={len(counts)}")
    failed_validations = Counter(
        (validation.severity, validation.check)
        for result in results
        for validation in result.validations
        if not validation.passed
    )
    for (severity, check), count in sorted(failed_validations.items()):
        add("Validation", f"{severity}: {check}", count)
    return rows


def load_baseline_rows(path: Path | None) -> dict[str, dict[str, Any]]:
    if path is None:
        return {}
    if not path.is_file():
        raise FileNotFoundError(f"baseline workbook not found: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Run_Metadata" not in workbook.sheetnames or "Cases" not in workbook.sheetnames:
        raise ValueError("baseline must contain Run_Metadata and Cases sheets")
    metadata_values = {
        row[0]: row[1]
        for row in workbook["Run_Metadata"].iter_rows(
            min_row=2, values_only=True
        )
        if row[0]
    }
    if str(metadata_values.get("Schema Version", "")) != SCHEMA_VERSION:
        raise ValueError(
            f"baseline schema {metadata_values.get('Schema Version')!r} "
            f"does not match {SCHEMA_VERSION!r}"
        )
    sheet = workbook["Cases"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    required_headers = {
        "ID",
        "Status",
        "Final Waypoints (target->entry)",
        "Fuzzer target_pcs (entry->target)",
    }
    missing_headers = required_headers - set(headers)
    if missing_headers:
        raise ValueError(f"baseline Cases sheet is missing columns: {missing_headers}")
    rows = {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        rows[str(row.get("ID", ""))] = row
    return rows


def build_change_rows(
    case_rows: list[dict[str, Any]],
    baseline_rows: dict[str, dict[str, Any]],
    include_removed: bool = False,
) -> list[dict[str, Any]]:
    fields = ["Status", "Failed Stage"]
    for _, label, _ in STAGES:
        fields.extend(stage_case_columns(label))
    fields.extend([
        "Final Listed PCs (target->entry)",
        "Final Resolved Targets (target->entry)",
        "Fuzzer target_pcs (entry->target)",
        "Bug Position Exact Relation",
        "Bug Position Function Relation",
        "Bug Position Instrumentation Target",
        "Bug Position PC32",
        "Bug Position Resolution Class",
        "Validation Failure Signature",
    ])
    changes: list[dict[str, Any]] = []
    if not baseline_rows:
        return [
            {
                "ID": "",
                "Field": "",
                "Baseline": "",
                "Current": "",
                "Change": "No baseline workbook supplied",
            }
        ]
    for current in case_rows:
        case_id = str(current["ID"])
        baseline = baseline_rows.get(case_id)
        if baseline is None:
            changes.append(
                {
                    "ID": case_id,
                    "Field": "Case",
                    "Baseline": "",
                    "Current": "present",
                    "Change": "added",
                }
            )
            continue
        for field_name in fields:
            old_value = baseline.get(field_name, "") or ""
            new_value = current.get(field_name, "") or ""
            if old_value != new_value:
                changes.append(
                    {
                        "ID": case_id,
                        "Field": field_name,
                        "Baseline": old_value,
                        "Current": new_value,
                        "Change": "modified",
                    }
                )
    if include_removed:
        current_ids = {str(row["ID"]) for row in case_rows}
        for case_id in sorted(
            set(baseline_rows) - current_ids, key=lambda value: int(value)
        ):
            changes.append(
                {
                    "ID": case_id,
                    "Field": "Case",
                    "Baseline": "present",
                    "Current": "",
                    "Change": "removed",
                }
            )
    return changes


def build_workbook(
    results: list[CaseResult],
    run_metadata: dict[str, Any],
    baseline_rows: dict[str, dict[str, Any]],
    oracle_entries: dict[str, dict[str, Any]],
    include_removed_baseline_cases: bool = False,
) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)

    append_mapping_sheet(
        workbook,
        "Run_Metadata",
        [{"Key": key, "Value": value} for key, value in run_metadata.items()],
    )

    case_rows = [case_row(result) for result in results]
    case_headers = list(BASE_CASE_COLUMNS)
    for _, label, _ in STAGES:
        case_headers.extend(stage_case_columns(label))
    case_headers.extend(FINAL_CASE_COLUMNS)
    cases_sheet = workbook.create_sheet("Cases")
    cases_sheet.append(case_headers)
    for row in case_rows:
        cases_sheet.append([row.get(header, "") for header in case_headers])
    style_sheet(cases_sheet)

    next_stage = {
        STAGES[index][0]: STAGES[index + 1][0]
        for index in range(len(STAGES) - 1)
    }
    long_rows: list[dict[str, Any]] = []
    for result in results:
        for stage_order, (stage_key, stage_label, stage_kind) in enumerate(STAGES):
            snapshot = result.snapshots.get(stage_key)
            if snapshot is None:
                continue
            following = result.snapshots.get(next_stage.get(stage_key, ""))
            following_ids = {
                node.identity for node in following.nodes
            } if following else set()
            for target_index, node in enumerate(snapshot.nodes):
                present_next = (
                    node.identity in following_ids if following is not None else ""
                )
                long_rows.append(
                    {
                        "Run ID": run_metadata["Run ID"],
                        "Case ID": result.metadata.get("ID", ""),
                        "Stage": stage_key,
                        "Stage Label": stage_label,
                        "Stage Kind": stage_kind,
                        "Stage Order": stage_order,
                        "Index target->entry": target_index,
                        "Index entry->target": len(snapshot.nodes) - target_index - 1,
                        "Function": node.func_name,
                        "Report Location": node.location,
                        "Report Target": node.target,
                        "Resolved Target": snapshot.resolved_targets[target_index],
                        "PC64": snapshot.pcs64[target_index],
                        "PC32": snapshot.pcs32[target_index],
                        "PC Resolution Error": snapshot.resolution_errors[target_index],
                        "Inline": node.is_inline,
                        "Original Index": node.original_index,
                        "BB Offset": node.bb_offset,
                        "BB Count": node.bb_count,
                        "Hot Entry Score": node.hot_entry_score,
                        "Value Score": node.value_score,
                        "Present Next Stage": present_next,
                        "Drop Transition": (
                            ""
                            if present_next in (True, "")
                            else f"removed_before_{next_stage[stage_key]}"
                        ),
                        "Causal Phase": "",
                    }
                )
    append_mapping_sheet(workbook, "Waypoints_Long", long_rows)

    validation_rows = [
        {
            "Run ID": run_metadata["Run ID"],
            "Case ID": result.metadata.get("ID", ""),
            "Check": validation.check,
            "Severity": validation.severity,
            "Passed": validation.passed,
            "Details": validation.details,
        }
        for result in results
        for validation in result.validations
    ]
    append_mapping_sheet(workbook, "Validation", validation_rows)
    append_mapping_sheet(workbook, "Summary", build_summary_rows(results))

    change_rows = build_change_rows(
        case_rows, baseline_rows, include_removed_baseline_cases
    )
    append_mapping_sheet(
        workbook, "Changes", change_rows
    )

    changed_case_ids = {
        str(row["ID"]) for row in change_rows if row.get("ID")
    }
    manual_rows = []
    for result in results:
        error_count = sum(
            not check.passed and check.severity == "error"
            for check in result.validations
        )
        if error_count:
            priority = "required: validation error"
        elif result.metadata.get("ID", "") in changed_case_ids:
            priority = "required: baseline change"
        elif len(result.report_sections) > 1:
            priority = "required: multi-section report"
        else:
            priority = "sample"
        oracle = oracle_entries.get(str(result.metadata.get("ID", "")), {})
        manual_rows.append(
            {
                "Case ID": result.metadata.get("ID", ""),
                "Title": result.metadata.get("Title", ""),
                "Priority": priority,
                "Expected Target": oracle.get("expected_target", ""),
                "Allowed Fallback Targets": json_list(
                    oracle.get("allowed_fallback_targets", [])
                ),
                "Required Nodes": json_list(oracle.get("required_nodes", [])),
                "Allowed Nodes": json_list(oracle.get("allowed_nodes", [])),
                "Forbidden Nodes": json_list(oracle.get("forbidden_nodes", [])),
                "Forbidden Report Sections": json_list(
                    oracle.get("forbidden_report_sections", [])
                ),
                "Target Verdict": oracle.get("target_verdict", ""),
                "Section Verdict": oracle.get("section_verdict", ""),
                "Chain Verdict": oracle.get("chain_verdict", ""),
                "Reviewer": oracle.get("reviewer", ""),
                "Review Date": oracle.get("review_date", ""),
                "Notes": oracle.get("notes", ""),
                "Accepted Exception": oracle.get("accepted_exception", ""),
            }
        )
    append_mapping_sheet(workbook, "Manual_Audit", manual_rows)

    legend_rows = [
        {
            "Item": "Chain direction",
            "Definition": (
                "All displayed waypoint chains and Listed PCs use target->entry order. "
                "Fuzzer target_pcs is explicitly entry->target."
            ),
        },
        {
            "Item": "Correctness",
            "Definition": (
                "Validation checks establish invariants, not a unique optimal chain. "
                "Target, report-section, and chain verdicts require the Manual_Audit sheet."
            ),
        },
        {
            "Item": "Zero PC",
            "Definition": (
                "0x00000000 marks a failed source-to-instrumentation resolution and is an error."
            ),
        },
        {
            "Item": "Blank raw-stage PC",
            "Definition": (
                "A null PC means the node was removed before the extractor's core "
                "instrumentation build, so resolution was not attempted."
            ),
        },
        {
            "Item": "Bug Position relation",
            "Definition": (
                "target means index 0; toward_entry means a later displayed index; absent means "
                "the exact location or enclosing function is not in the final chain."
            ),
        },
        {
            "Item": "Outlier removal",
            "Definition": (
                "This engineering phase is retained outside the four-rule paper pipeline."
            ),
        },
    ]
    for key, label, kind in STAGES:
        legend_rows.append(
            {"Item": f"Stage {key}", "Definition": f"{label} ({kind})."}
        )
    append_mapping_sheet(workbook, "Legend", legend_rows)
    return workbook


def save_workbook_atomic(workbook: Workbook, destination: Path) -> None:
    temporary = destination.with_name(
        f".{destination.stem}.writing{destination.suffix}"
    )
    workbook.save(temporary)
    os.replace(temporary, destination)


def canonical_pc(value: str | None, zero_value: str) -> str | None:
    return None if value in (None, zero_value) else value


def snapshot_payload(snapshot: StageSnapshot | None) -> dict[str, Any] | None:
    if snapshot is None:
        return None
    nodes = []
    for index, node in enumerate(snapshot.nodes):
        nodes.append(
            {
                **asdict(node),
                "waypoint": node.target,
                "resolved_target": snapshot.resolved_targets[index],
                "pc64": canonical_pc(snapshot.pcs64[index], ZERO_PC64),
                "pc32": canonical_pc(snapshot.pcs32[index], ZERO_PC32),
                "resolution_error": snapshot.resolution_errors[index] or None,
            }
        )
    return {
        "waypoints_target_to_entry": snapshot.targets,
        "resolved_targets_target_to_entry": snapshot.resolved_targets,
        "pcs64_target_to_entry": [
            canonical_pc(value, ZERO_PC64) for value in snapshot.pcs64
        ],
        "pcs32_target_to_entry": [
            canonical_pc(value, ZERO_PC32) for value in snapshot.pcs32
        ],
        "resolution_errors_target_to_entry": [
            value or None for value in snapshot.resolution_errors
        ],
        "nodes": nodes,
    }


def referenced_source_path(kernel_dir: Path, value: str) -> str | None:
    """Return a kernel-relative source path from a location or resolved target."""
    location = value.rsplit("@", 1)[-1].strip()
    if ":" not in location:
        return None
    path_text, line = location.rsplit(":", 1)
    if not line.isdigit() or not path_text:
        return None
    path = Path(path_text)
    if path.is_absolute():
        try:
            path = path.resolve().relative_to(kernel_dir.resolve())
        except ValueError:
            return None
    return path.as_posix()


def case_input_state(result: CaseResult) -> dict[str, Any]:
    """Capture only inputs that can affect this case's extracted chain or PCs."""
    kernel_dir = Path(result.metadata.get("Kernel Dir", ""))
    source_paths = {
        source_path
        for snapshot in result.snapshots.values()
        for node in snapshot.nodes
        if (source_path := referenced_source_path(kernel_dir, node.location))
    }
    for value in (
        str(result.metadata.get("Bug Position", "")),
        result.bug_position_source_target,
        result.bug_position_instrumentation_target,
    ):
        source_path = referenced_source_path(kernel_dir, value)
        if source_path:
            source_paths.add(source_path)
    return {
        "title": file_state(Path(result.metadata.get("Title Path", ""))),
        "report": file_state(Path(result.metadata.get("Report Path", ""))),
        "bzimage": file_state(kernel_dir / "arch/x86/boot/bzImage"),
        "vmlinux_identity": get_vmlinux_cache_identity(str(kernel_dir)),
        "referenced_sources": {
            relative: file_state(kernel_dir / relative)
            for relative in sorted(source_paths)
        },
    }


def build_canonical_payload(
    results: list[CaseResult], run_metadata: dict[str, Any]
) -> dict[str, Any]:
    """Build the machine-readable source of truth behind the XLSX view."""
    cases = []
    for result in results:
        stage_payload = {
            key: {
                "label": label,
                "kind": kind,
                "chain": snapshot_payload(result.snapshots.get(key)),
            }
            for key, label, kind in STAGES
        }
        final = snapshot_payload(
            operational_final_snapshot(result.snapshots.get("outlier_removal"))
        )
        if final is not None:
            final["fuzzer_pcs32_entry_to_target"] = list(
                reversed(final["pcs32_target_to_entry"])
            )
        cases.append(
            {
                "case_id": str(result.metadata.get("ID", "")),
                "title": result.metadata.get("Title", ""),
                "bug_position": result.metadata.get("Bug Position", ""),
                "status": result.status,
                "failed_stage": result.failed_stage or None,
                "error": result.error or None,
                "warnings": result.warnings,
                "report_type": result.report_type,
                "report_sections": result.report_sections,
                "paths": {
                    "kernel_dir": result.metadata.get("Kernel Dir", ""),
                    "title": result.metadata.get("Title Path", ""),
                    "report": result.metadata.get("Report Path", ""),
                },
                "input_state": case_input_state(result),
                "bug_position_resolution": {
                    "function": result.bug_position_func or None,
                    "source_target": result.bug_position_source_target or None,
                    "instrumentation_target": (
                        result.bug_position_instrumentation_target or None
                    ),
                    "pc64": canonical_pc(result.bug_position_pc64, ZERO_PC64),
                    "pc32": canonical_pc(result.bug_position_pc32, ZERO_PC32),
                    "resolution_class": result.bug_position_resolution_class,
                    "exact_relation": result.bug_position_exact_relation,
                    "exact_index": result.bug_position_exact_index,
                    "function_relation": result.bug_position_function_relation,
                    "function_index": result.bug_position_function_index,
                },
                "stages": stage_payload,
                "final": final,
                "validations": [asdict(item) for item in result.validations],
            }
        )
    return {
        "schema_version": CANONICAL_SCHEMA_VERSION,
        "artifact_type": "script_waypoint_extraction",
        "chain_direction": CHAIN_DIRECTION,
        "run_metadata": dict(run_metadata),
        "stage_definitions": [
            {"key": key, "label": label, "kind": kind}
            for key, label, kind in STAGES
        ],
        "cases": cases,
    }


def write_json_atomic(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.writing")
    temporary.write_text(
        json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    os.replace(temporary, path)


def parse_case_ids(value: str | None) -> set[str] | None:
    if not value:
        return None
    return {part.strip() for part in value.split(",") if part.strip()}


def default_output_path(run_id: str) -> Path:
    return (
        PROJECT_ROOT
        / "agent_analysis"
        / "waypoints_regression"
        / run_id
        / "waypoints_regression.xlsx"
    )


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Run phase-level waypoint extraction and write an immutable XLSX regression artifact."
    )
    parser.add_argument(
        "--benchmark",
        type=Path,
        default=PROJECT_ROOT / "benchmark" / "benchmark.csv",
    )
    parser.add_argument(
        "--cases-root",
        type=Path,
        default=Path.home() / "kernels" / "SyzPilot-experiments" / "cases",
    )
    parser.add_argument(
        "--configs-dir",
        type=Path,
        default=Path.home() / "kernels" / "SyzPilot-experiments" / "configs",
    )
    parser.add_argument("--output", type=Path)
    parser.add_argument(
        "--json-output",
        type=Path,
        help="Machine-readable output (default: output path with .json suffix).",
    )
    parser.add_argument("--baseline", type=Path)
    parser.add_argument(
        "--oracle",
        type=Path,
        default=PROJECT_ROOT / "benchmark" / "waypoints_oracle.json",
    )
    parser.add_argument(
        "--case-ids", help="Optional comma-separated benchmark IDs for a targeted run."
    )
    parser.add_argument(
        "--checkpoint-every",
        type=int,
        default=5,
        help="Write a partial workbook after this many completed cases.",
    )
    args = parser.parse_args()
    if args.checkpoint_every < 1:
        parser.error("--checkpoint-every must be at least 1")

    try:
        baseline_rows = load_baseline_rows(args.baseline)
        oracle_entries = load_oracle(args.oracle)
    except Exception as exc:
        parser.error(str(exc))

    git_sha, git_dirty = get_git_state()
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    run_id = f"{timestamp}_{git_sha[:12]}_{uuid.uuid4().hex[:12]}"
    output_path = args.output or default_output_path(run_id)
    json_output_path = args.json_output or output_path.with_suffix(".json")
    partial_path = output_path.with_suffix(".partial.xlsx")
    checkpoint_writing_path = partial_path.with_name(
        f".{partial_path.stem}.writing{partial_path.suffix}"
    )
    artifact_paths = (
        output_path,
        partial_path,
        checkpoint_writing_path,
        json_output_path,
        json_output_path.with_name(f".{json_output_path.name}.writing"),
    )
    resolved_artifact_paths = [path.resolve() for path in artifact_paths]
    if len(resolved_artifact_paths) != len(set(resolved_artifact_paths)):
        parser.error("output, JSON, partial, and temporary paths must be distinct")
    for candidate in artifact_paths:
        if candidate.exists():
            parser.error(f"output already exists and will not be overwritten: {candidate}")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    selected_ids = parse_case_ids(args.case_ids)
    benchmark_rows = load_benchmark(args.benchmark)
    if selected_ids is not None:
        benchmark_rows = [row for row in benchmark_rows if row["ID"] in selected_ids]

    run_metadata = {
        "Schema Version": SCHEMA_VERSION,
        "Run ID": run_id,
        "Run State": "starting",
        "UTC Timestamp": timestamp,
        "Git SHA": git_sha,
        "Git Dirty": git_dirty,
        "Benchmark Path": str(args.benchmark.resolve()),
        "Cases Root": str(args.cases_root.resolve()),
        "Configs Dir": str(args.configs_dir.resolve()),
        "Oracle Path": str(args.oracle.resolve()) if args.oracle else "",
        "Output Path": str(output_path.resolve()),
        "Baseline Path": str(args.baseline.resolve()) if args.baseline else "",
        "Chain Direction": CHAIN_DIRECTION,
        "Python": sys.version.replace("\n", " "),
        "Platform": platform.platform(),
        "Command": " ".join(sys.argv),
        "Selected Case IDs": ",".join(sorted(selected_ids)) if selected_ids else "all",
        "B_min": WaypointsExtractor.MIN_BB_COUNT,
        "K_min": WaypointsExtractor.MIN_WAYPOINTS_LENGTH,
        "K_max": WaypointsExtractor.MAX_WAYPOINTS_LENGTH,
        "Hot Entry Threshold": WaypointsExtractor.HOT_ENTRY_SCORE_THRESHOLD,
        "Value Score Alpha": WaypointsExtractor.VALUE_SCORE_ALPHA,
        "Value Score Beta": WaypointsExtractor.VALUE_SCORE_BETA,
        "Value Score Gamma": WaypointsExtractor.VALUE_SCORE_GAMMA,
    }

    results = []
    for index, row in enumerate(benchmark_rows, start=1):
        print(
            f"[{index}/{len(benchmark_rows)}] case_{row['ID']}: {row['Title']}",
            flush=True,
        )
        result = run_case(
            row,
            args.cases_root,
            args.configs_dir,
            oracle_entries.get(str(row["ID"])),
        )
        results.append(result)
        print(
            f"  status={result.status} failed_stage={result.failed_stage or '-'}",
            flush=True,
        )
        if index % args.checkpoint_every == 0:
            run_metadata["Run State"] = f"in_progress:{index}/{len(benchmark_rows)}"
            workbook = build_workbook(
                results,
                run_metadata,
                baseline_rows,
                oracle_entries,
                include_removed_baseline_cases=False,
            )
            save_workbook_atomic(workbook, partial_path)
            write_json_atomic(
                json_output_path,
                build_canonical_payload(results, run_metadata),
            )

    run_metadata["Run State"] = f"complete:{len(results)}/{len(benchmark_rows)}"
    workbook = build_workbook(
        results,
        run_metadata,
        baseline_rows,
        oracle_entries,
        include_removed_baseline_cases=selected_ids is None,
    )
    save_workbook_atomic(workbook, partial_path)
    os.replace(partial_path, output_path)
    write_json_atomic(
        json_output_path,
        build_canonical_payload(results, run_metadata),
    )
    print(f"Wrote {output_path}", flush=True)
    print(f"Wrote {json_output_path}", flush=True)
    validation_error_count = sum(
        not validation.passed and validation.severity == "error"
        for result in results
        for validation in result.validations
    )
    if validation_error_count:
        print(
            f"Regression completed with {validation_error_count} validation errors.",
            flush=True,
        )
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
